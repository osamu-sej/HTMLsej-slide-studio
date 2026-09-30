"""tools/import_deck.py: nothing is lost or garbled, and the page's structure comes back."""
import io
import sys
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))

import import_deck  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402


def text(slide, x, y, w, h, lines, size=20, bold=False):
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        run = (frame.paragraphs[0] if i == 0 else frame.add_paragraph()).add_run()
        run.text, run.font.size, run.font.bold = line, Pt(size), bold


def deck() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[5]  # title only
    cover = prs.slides.add_slide(layout)
    cover.shapes.title.text = "改革の報告"
    cards = prs.slides.add_slide(layout)
    cards.shapes.title.text = "3つの成果"
    for i, (head, desc) in enumerate([("成果", "作業時間が42%減った"), ("要因", "手順を写真で標準化した"), ("課題", "定着に時間がかかる店がある")]):
        text(cards, 0.5 + i * 4, 2.0, 3.8, 0.6, head, 24, True)
        text(cards, 0.5 + i * 4, 2.65, 3.8, 1.5, desc, 20)
    text(cards, 0.5, 4.4, 11.8, 0.6, "次のアクション：全店展開を決める", 20, True)
    kpi = prs.slides.add_slide(layout)
    kpi.shapes.title.text = "主要な数値"
    for i, (value, label) in enumerate([("-120h/月", "作業時間"), ("-0.3pt", "廃棄率")]):
        text(kpi, 0.5 + i * 4, 2.0, 3.5, 0.7, value, 32, True)
        text(kpi, 0.5 + i * 4, 2.75, 3.5, 0.5, label, 20)
    compare = prs.slides.add_slide(layout)
    compare.shapes.title.text = "改革の前と後"
    text(compare, 0.5, 2.0, 5.5, 0.6, "改革前", 24, True)
    text(compare, 0.5, 2.65, 5.5, 2.0, ["紙で発注していた", "検品は目視だった"], 20)
    text(compare, 6.8, 2.0, 5.5, 0.6, "改革後", 24, True)
    text(compare, 6.8, 2.65, 5.5, 2.0, ["端末で発注する", "検品は端末で行う"], 20)
    long = prs.slides.add_slide(layout)
    long.shapes.title.text = "やること一覧"
    text(long, 0.5, 1.5, 12, 5.5, [f"作業{i}を見直す" for i in range(13)], 18)
    grouped = prs.slides.add_slide(layout)
    grouped.shapes.title.text = "グループ"
    group = grouped.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(8), Inches(6), Inches(3), Inches(1)).text_frame.text = "下の文"
    group.shapes.add_textbox(Inches(8), Inches(1), Inches(3), Inches(1)).text_frame.text = "上の文"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


class PowerPointTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.title, cls.slides, cls.stats = import_deck.read_pptx(deck())

    def test_cards_read_heading_with_description_and_footer_stays_apart(self):
        self.assertEqual(self.slides[1]["points"], ["成果：作業時間が42%減った", "要因：手順を写真で標準化した",
                                                    "課題：定着に時間がかかる店がある", "次のアクション：全店展開を決める"])

    def test_kpi_cards_come_back_as_kpi_with_their_sign(self):
        self.assertEqual(self.slides[2]["type"], "kpi")
        self.assertEqual([(i["label"], i["value"]) for i in self.slides[2]["items"]], [("作業時間", "-120h/月"), ("廃棄率", "-0.3pt")])

    def test_two_sides_come_back_as_compare(self):
        s = self.slides[3]
        self.assertEqual((s["type"], s["leftTitle"], s["rightTitle"]), ("compare", "改革前", "改革後"))
        self.assertEqual(s["rightItems"], ["端末で発注する", "検品は端末で行う"])

    def test_long_lists_continue_on_a_new_slide_instead_of_being_cut(self):
        points = [p for s in self.slides if s["title"].startswith("やること一覧") for p in s["points"]]
        self.assertEqual(len(points), 13)
        self.assertEqual(self.slides[5]["title"], "やること一覧（続き）")

    def test_group_members_are_read_in_slide_order(self):
        self.assertEqual(self.slides[6]["points"], ["上の文", "下の文"])

    def test_minus_signs_and_bullets(self):
        self.assertEqual(import_deck.clean_line("-42%"), "-42%")
        self.assertEqual(import_deck.clean_line("- 手順を見直す"), "手順を見直す")
        self.assertEqual(import_deck.clean_line("・手順を見直す"), "手順を見直す")


def docx(body: str) -> bytes:
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    styles = (f'<w:styles {ns}><w:style w:type="paragraph" w:styleId="a3"><w:name w:val="Title"/></w:style>'
              '<w:style w:type="paragraph" w:styleId="10"><w:name w:val="heading 1"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>'
              '<w:style w:type="paragraph" w:styleId="20"><w:name w:val="heading 2"/><w:pPr><w:outlineLvl w:val="1"/></w:pPr></w:style>'
              '<w:style w:type="paragraph" w:styleId="mine"><w:name w:val="My Heading"/><w:basedOn w:val="20"/></w:style></w:styles>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/styles.xml", styles)
        z.writestr("word/document.xml", f"<w:document {ns}><w:body>{body}</w:body></w:document>")
    return out.getvalue()


def p(text, style=None, bold=False):
    ppr = '<w:pPr><w:pStyle w:val="%s"/></w:pPr>' % style if style else ""
    rpr = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f"<w:p>{ppr}<w:r>{rpr}<w:t>{text}</w:t></w:r></w:p>"


class WordTest(unittest.TestCase):
    def test_heading_styles_by_outline_level_tables_and_chapters(self):
        table = "<w:tbl>" + "".join("<w:tr>" + "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in row) + "</w:tr>"
                                    for row in [["指標", "前", "後"], ["作業時間", "95分", "54分"]]) + "</w:tbl>"
        title, slides, _ = import_deck.read_docx(docx(p("報告", "a3") + p("第1章", "10") + p("現状", "20") + p("人手不足。") + p("効果", "mine") + table))
        self.assertEqual(title, "報告")
        self.assertEqual([s["type"] for s in slides], ["title", "section", "content", "table"])
        self.assertEqual(slides[3]["rows"], [["作業時間", "95分", "54分"]])

    def test_bold_lines_are_headings_when_there_are_no_heading_styles(self):
        _, slides, _ = import_deck.read_docx(docx(p("報告書") + p("はじめに", bold=True) + p("目的を書く。") + p("結果", bold=True) + p("効果が出た。")))
        self.assertEqual([s["title"] for s in slides], ["報告書", "はじめに", "結果"])


class ChromeTextTest(unittest.TestCase):
    def test_template_furniture_is_dropped_but_content_is_kept(self):
        for line in ["社内限り", "Confidential", "© 2026 Example Inc.", "12"]:
            self.assertTrue(import_deck.CHROME_TEXT.match(line), line)
        for line in ["売上は前年比108%", "次のアクション"]:
            self.assertFalse(import_deck.CHROME_TEXT.match(line), line)


class PdfTest(unittest.TestCase):
    def test_headers_page_numbers_titles_and_wrapped_paragraphs(self):
        import fitz

        doc = fitz.open()
        for n in range(3):
            page = doc.new_page(width=960, height=540)
            page.insert_text((40, 30), "株式会社テスト", fontname="japan", fontsize=10)
            page.insert_text((40, 90), f"見出し{n + 1}", fontname="japan", fontsize=32)
            page.insert_text((40, 160), "一つ目の段落はとても長いので途中で", fontname="japan", fontsize=20)
            page.insert_text((40, 186), "折り返している。", fontname="japan", fontsize=20)
            page.insert_text((40, 230), "二つ目の段落。", fontname="japan", fontsize=20)
            page.insert_text((900, 520), str(n + 1), fontname="japan", fontsize=10)
        title, slides, _ = import_deck.read_pdf(doc.tobytes())
        self.assertEqual(title, "見出し1")
        self.assertEqual(slides[1]["title"], "見出し2")
        # Wrapped lines join into one paragraph; a sentence right under the title is the slide's message.
        self.assertEqual(slides[1]["takeaway"], "一つ目の段落はとても長いので途中で折り返している。")
        self.assertEqual(slides[1]["points"], ["二つ目の段落。"], "the repeated header and the page number are gone")


if __name__ == "__main__":
    unittest.main()
