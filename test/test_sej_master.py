"""The SEJ theme draws the template's masters where assets/sej/template.pptx puts them (engine.js SEJ_MASTER/SEJ_BOX)."""
import hashlib
import json
import re
import unittest
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "assets" / "sej" / "template.pptx"
ENGINE = (ROOT / "public" / "engine" / "engine.js").read_text("utf-8")
PX = 144 / 914400  # 13.333 in → 1920 px


def js_object(name: str) -> dict:
    """The literal `const NAME = {...};` from engine.js as a Python dict (numbers and strings only)."""
    body = re.search(rf"const {name} = (\{{.*?\}});\n", ENGINE, re.S).group(1)
    body = re.sub(r"(\w+):", r'"\1":', body)
    return json.loads(re.sub(r",\s*([}\]])", r"\1", body))


def boxes(master):
    """Non-placeholder shapes of a master as {name-ish key: [x, y, w, h] in slide px}."""
    found = {}
    for sh in master.shapes:
        if sh.is_placeholder:
            continue
        box = [sh.left * PX, sh.top * PX, sh.width * PX, sh.height * PX]
        text = sh.text_frame.text.strip() if sh.has_text_frame else ""
        if sh.shape_type == 13:
            key = "logo" if sh.width < sh.height * 2 else "copyright" if sh.width < 3000000 else "rule"
        else:
            key = {"秘（B）": "secret", "社内限り": "internal", "‹#›": "page"}.get(text, "slogan" if "笑顔" in text else text)
        found[key] = box
    return found


class SejMasterTest(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation(TEMPLATE)
        self.master = js_object("SEJ_MASTER")
        self.box = js_object("SEJ_BOX")

    def near(self, got, want, what):
        for g, w in zip(got, want):
            self.assertAlmostEqual(g, w, delta=1.0, msg=f"{what}: {got} vs template {want}")

    def test_template_is_the_reviewed_file(self):
        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
                         json.loads((ROOT / "assets" / "sej" / "brand-profile.json").read_text("utf-8"))["templateSha256"])
        self.assertEqual((self.prs.slide_width, self.prs.slide_height), (12192000, 6858000))

    def test_masters_match_the_template(self):
        for key, master in (("title", self.prs.slide_masters[0]), ("content", self.prs.slide_masters[1])):
            found = boxes(master)
            # The green rule is a picture of a solid band: rows 24–45 of its 69 are green.
            x, y, w, h = found["rule"]
            self.near(self.master[key]["rule"], [x, y + h * 24 / 69, w, h * 22 / 69], f"{key} rule")
            self.near(self.master[key]["slogan"][:4], found["slogan"], f"{key} slogan")
            for part in ("logo", "secret", "internal", "copyright"):
                self.near(self.box[part], found[part], f"{key} {part}")
            self.assertEqual("page" in found, self.master[key]["page"], f"{key} page number")
            if "page" in found:
                self.near(self.box["page"], found["page"], f"{key} page")


if __name__ == "__main__":
    unittest.main()
