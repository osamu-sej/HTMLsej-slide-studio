"""Source pages retain their own masters, objects, text, charts, effects, and hidden slides."""
import io
import sys
import unittest
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "tools"))
sys.path.insert(0, str(HERE))

import pptx_exact  # noqa: E402
import pptx_fixtures  # noqa: E402
from pptx_exact import read_pptx_exact  # noqa: E402

IN = 144  # slide pixels per inch


def by_text(slide, words):
    return next(o for o in slide["elements"] if words in o.get("text", ""))


class SejDeckTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.deck = read_pptx_exact(pptx_fixtures.sej_deck())
        cls.slides = cls.deck["slideData"]

    def test_an_sej_deck_keeps_every_slide_including_hidden_ones(self):
        self.assertTrue(self.deck["sej"])
        self.assertEqual(self.deck["fidelity"], "exact")
        self.assertEqual(self.deck["deckTitle"], "発注業務の改革")
        self.assertEqual(len(self.slides), 4)
        self.assertEqual(self.deck["stats"]["hidden"], 1)
        self.assertIn("隠したページ", str(self.slides[2]))
        self.assertTrue(self.slides[2]["hidden"])
        for slide in self.slides:
            self.assertEqual(slide["type"], "blank")
            self.assertTrue(slide["hideTitle"], "the slide's own title box shows the title")
        self.assertEqual([s["master"] for s in self.slides], ["source"] * 4)

    def test_the_original_master_marks_are_editable_objects(self):
        words = " ".join(o.get("text", "") for s in self.slides for o in s["elements"])
        for mark in ("社内限り", "明日の笑顔", "秘"):
            self.assertIn(mark, words)
        self.assertGreater(self.deck["stats"]["pictures"], 0, "the original logo is copied")

    def test_placeholders_take_their_place_from_the_layout(self):
        title = by_text(self.slides[1], "改革の全体像")
        self.assertLess(title["y"], 130)
        self.assertLess(title["x"], 120)
        self.assertTrue(title.get("bold"))
        body = by_text(self.slides[3], "最初の点")
        self.assertIn('data-indent="1"', body["text"], "the second level is indented")

    def test_shapes_keep_their_place_fill_and_text_runs(self):
        box = by_text(self.slides[1], "現状の課題")
        self.assertEqual((box["kind"], box["shape"], box["fill"], box["stroke"]), ("shape", "roundRect", "#dce4f2", "none"))
        for key, inches in (("x", 0.6), ("y", 1.6), ("w", 3.6), ("h", 1.6)):
            self.assertAlmostEqual(box[key], inches * IN, delta=1, msg=key)
        self.assertEqual(box["color"], "#1f3864")
        self.assertTrue(box["bold"])
        self.assertIn("font-size: 40.0px", box["text"], "20pt is 40px")
        self.assertIn("color: #1a1a1a; font-weight: 400", box["text"], "the second run is black and regular")
        oval = by_text(self.slides[1], "30分")
        self.assertEqual((oval["shape"], oval["fill"], oval["fs"]), ("ellipse", "#f5f0ea", 56.0))

    def test_bullets_and_links(self):
        text = by_text(self.slides[1], "端末で発注する")["text"]
        self.assertEqual(text.count('data-bullet="■"'), 3)
        self.assertIn('<a href="https://example.com/portal">', text)

    def test_tables_cell_by_cell(self):
        table = next(o for o in self.slides[1]["elements"] if o["kind"] == "table")
        self.assertEqual(len(table["cells"]), 3)
        self.assertEqual([c["text"] for c in table["cells"][0]], ["<p>項目</p>", "<p>現状</p>", "<p>目標</p>"])
        self.assertEqual(table["cells"][2][2]["text"], "<p>2.5%</p>")
        self.assertTrue(table["cells"][0][0].get("fill"), "the header row is filled from the table style")
        self.assertAlmostEqual(sum(table["cols"]), 1, places=3)

    def test_charts_keep_their_data_and_formatting(self):
        chart = next(o for o in self.slides[1]["elements"] if o["kind"] == "chart")["chart"]
        self.assertEqual(chart["type"], "bar")
        self.assertEqual(chart["labels"], ["A店", "B店", "C店"])
        self.assertEqual(chart["series"][0]["values"], [42.0, 30.0, 18.0])
        st = chart["style"]
        self.assertEqual((st["dir"], st["gap"]), ("bar", 60.0))
        self.assertEqual(st["series"][0]["color"], "#b7c3da")
        self.assertEqual(st["series"][0]["points"], [{"i": 0, "color": "#1f3864"}])
        self.assertEqual(st["series"][0]["label"]["format"], '0"h"')
        self.assertEqual(st["series"][0]["label"]["pos"], "outEnd")
        self.assertTrue(st["val"]["hide"])
        self.assertEqual(st["font"]["size"], 28.0, "14pt")

    def test_lines_groups_and_notes(self):
        slide = self.slides[1]
        line = next(o for o in slide["elements"] if o["kind"] == "line")
        self.assertEqual((line["stroke"], line["strokeW"], line["tail"]), ("#1f3864", 4.0, "triangle"))
        groups = {o.get("group") for o in slide["elements"] if "グループ" in o.get("text", "")}
        self.assertEqual(len(groups), 1)
        self.assertIsNotNone(groups.pop())
        self.assertEqual(slide["notes"], "ここで全体像を説明します。")

    def test_buttons_keep_where_they_go_in_a_slide_show(self):
        self.assertEqual([s["sid"] for s in self.slides], ["p1", "p2", "p3", "p4"], "ids for every slide")
        back = by_text(self.slides[3], "全体像へ戻る")
        self.assertEqual(back["action"], {"type": "slide", "to": "p2"})
        self.assertEqual(by_text(self.slides[3], "終わる")["action"], {"type": "end"})
        self.assertEqual(back.get("overAction"), {"type": "next"}, "マウスの通過")
        self.assertEqual((back.get("cap"), back.get("cmpd"), back.get("join")), ("round", "dbl", "bevel"), "線端・複合線・結合点")
        self.assertNotIn("cmpd", by_text(self.slides[3], "終わる"))
        end = by_text(self.slides[3], "終わる")
        self.assertEqual(end.get("glow"), {"r": 12, "color": "#b7c3da", "opacity": 0.6}, "光彩")
        self.assertEqual(end.get("soft"), 10, "ぼかし")
        self.assertEqual(end.get("reflect"), {"size": 0.3, "opacity": 0.4, "gap": 4}, "反射")
        self.assertNotIn("glow", back)
        self.assertNotIn("overAction", by_text(self.slides[3], "終わる"))

    def test_animations_and_the_transition(self):
        slide = self.slides[1]
        self.assertEqual(slide["transition"], "fade")
        first, second = slide["timeline"]
        self.assertEqual((first["el"], first["cls"], first["fx"], first["start"]), (by_text(slide, "現状の課題")["id"], "in", "fade", "click"))
        self.assertEqual((second["el"], second["fx"], second["dir"], second["start"]), (by_text(slide, "30分")["id"], "flyIn", "left", "after"))
        self.assertEqual(first["dur"], 500)


class PlainDeckTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.deck = read_pptx_exact(pptx_fixtures.plain_deck())
        cls.slides = cls.deck["slideData"]

    def test_a_4_3_deck_is_centred_on_the_wide_page(self):
        self.assertFalse(self.deck["sej"])
        box = by_text(self.slides[1], "現状の課題")
        self.assertAlmostEqual(box["x"], (1920 - 1440) / 2 + 0.6 * IN, delta=1)
        self.assertAlmostEqual(box["w"], 3.6 * IN, delta=1)

    def test_backgrounds_pictures_and_rotation(self):
        cover = self.slides[0]["elements"]
        self.assertEqual(cover[0]["fill"], "#f1f5fb", "the slide's background comes first")
        shapes = self.slides[2]["elements"]
        self.assertTrue(any(o["kind"] == "image" and o["src"].startswith("data:image/") for o in shapes))
        chevron = next(o for o in shapes if o.get("shape") == "chevron")
        self.assertEqual(chevron["rot"], 20)
        self.assertTrue(any(o.get("shape") == "star5" for o in shapes))

    def test_objects_off_the_slide_are_kept_for_editing(self):
        # The fixture puts a picture beyond the 4:3 page; the HTML viewport clips it.
        pictures = [o for o in self.slides[1]["elements"] if o["kind"] == "image"]
        self.assertEqual(len(pictures), 1)
        self.assertGreater(pictures[0]["x"], self.slides[1]["sourceViewport"]["x"] + self.slides[1]["sourceViewport"]["w"])

    def test_a_pie_chart_is_a_donut_without_a_hole_showing_percentages(self):
        chart = next(o for o in self.slides[2]["elements"] if o["kind"] == "chart")["chart"]
        self.assertEqual(chart["type"], "donut")
        self.assertEqual(chart["style"]["hole"], 0)
        self.assertTrue(chart["style"]["series"][0]["label"]["pct"])
        self.assertEqual(len(chart["style"]["series"][0]["points"]), 3, "each slice its own colour")


class ChartNumberTest(unittest.TestCase):
    def test_category_numbers_and_dates_as_the_chart_shows_them(self):
        self.assertEqual(pptx_exact.number_text(2024, "General"), "2024")
        self.assertEqual(pptx_exact.number_text(0.25, "0%"), "25%")
        self.assertEqual(pptx_exact.number_text(1234.5, "#,##0"), "1,235")
        self.assertEqual(pptx_exact.number_text(45566, "yyyy/m/d"), "2024/10/1")
        self.assertEqual(pptx_exact.number_text(45566, 'm"月"'), "10月")


class TransitionSoundTest(unittest.TestCase):
    """A slide's built-in transition sound comes over as one of the studio's synthesised sounds."""

    @staticmethod
    def _slide(sound_xml):
        ns = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        return etree.fromstring(f'<p:sld {ns}><p:transition spd="fast"><p:fade/>{sound_xml}</p:transition></p:sld>')

    def test_built_in_sounds_by_their_file_name(self):
        for name, kind in [("chime.wav", "chime"), ("applause.wav", "applause"), ("drumroll.wav", "drum"), ("breeze.wav", "whoosh"), ("cashreg.wav", "coin"), ("other.wav", "chime")]:
            out = pptx_exact.transition_of(self._slide(f'<p:sndAc><p:stSnd><p:snd r:embed="rId9" name="{name}"/></p:stSnd></p:sndAc>'))
            self.assertEqual(out["transitionSound"], kind, name)
            self.assertEqual(out["transition"], "fade")

    def test_transition_directions(self):
        def tr(effect):
            ns = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"'
            return pptx_exact.transition_of(etree.fromstring(f'<p:sld {ns}><p:transition>{effect}</p:transition></p:sld>'))
        self.assertEqual(tr('<p:push dir="u"/>'), {"transition": "push"}, "from below is the default")
        self.assertEqual(tr('<p:push dir="d"/>')["transitionDir"], "up")
        self.assertEqual(tr('<p:push/>')["transitionDir"], "right", "PowerPoint's default push comes from the right")
        self.assertEqual(tr('<p:wipe dir="r"/>')["transitionDir"], "left")
        self.assertNotIn("transitionDir", tr('<p:wipe/>'))
        self.assertEqual(tr('<p:cover dir="u"/>'), {"transition": "slide", "transitionDir": "down"})
        self.assertEqual(tr('<p:blinds dir="vert"/>')["transitionDir"], "vertical")
        self.assertEqual(tr('<p:split orient="horz"/>')["transitionDir"], "horizontal")
        self.assertNotIn("transitionDir", tr('<p:split orient="vert"/>'))
        self.assertNotIn("transitionDir", tr('<p:fade/>'))

    def test_stop_previous_sound_and_no_sound(self):
        self.assertEqual(pptx_exact.transition_of(self._slide("<p:sndAc><p:endSnd/></p:sndAc>"))["transitionSound"], "stop")
        self.assertNotIn("transitionSound", pptx_exact.transition_of(self._slide("")))


class SourceEffectsTest(unittest.TestCase):
    def test_every_slide_beyond_the_old_fifty_slide_limit_survives(self):
        prs = Presentation()
        for i in range(51):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = f"Page {i + 1}"
        out = io.BytesIO()
        prs.save(out)
        deck = read_pptx_exact(out.getvalue())
        self.assertEqual(len(deck["slideData"]), 51)
        self.assertIn("Page 51", str(deck["slideData"][-1]))

    def test_a_source_gradient_shadow_and_font_survive_import(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(4), Inches(2))
        shape.text = "Gradient"
        shape.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        sppr = shape._element.spPr
        for child in list(sppr):
            if etree.QName(child).localname in ("solidFill", "noFill"):
                sppr.remove(child)
        sppr.append(etree.fromstring('''<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:gsLst>
          <a:gs pos="0"><a:srgbClr val="FF0000"/></a:gs><a:gs pos="100000"><a:srgbClr val="0000FF"/></a:gs>
          </a:gsLst><a:lin ang="5400000" scaled="1"/></a:gradFill>'''))
        sppr.append(etree.fromstring('''<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
          <a:outerShdw blurRad="91440" dist="91440" dir="0"><a:srgbClr val="000000"><a:alpha val="50000"/></a:srgbClr></a:outerShdw>
          </a:effectLst>'''))
        out = io.BytesIO()
        prs.save(out)
        deck = read_pptx_exact(out.getvalue())
        imported = by_text(deck["slideData"][0], "Gradient")
        self.assertEqual(deck["slideData"][0]["master"], "source")
        self.assertEqual(imported["fontFace"], "Arial")
        self.assertEqual([s["color"] for s in imported["gradient"]["stops"]], ["#ff0000", "#0000ff"])
        self.assertEqual(imported["gradient"]["angle"], 90)
        self.assertAlmostEqual(imported["shadow"]["opacity"], 0.5)

    def test_text_shadow_and_glow_on_the_letters_become_text_effects(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text = "光る文字"
        rpr = box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
        rpr.append(etree.fromstring('''<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="19050"><a:solidFill><a:srgbClr val="1F3864"/></a:solidFill></a:ln>'''))
        rpr.append(etree.fromstring('''<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
          <a:glow rad="76200"><a:srgbClr val="B7C3DA"><a:alpha val="60000"/></a:srgbClr></a:glow>
          <a:outerShdw blurRad="38100" dist="38100" dir="2700000"><a:srgbClr val="000000"><a:alpha val="40000"/></a:srgbClr></a:outerShdw>
          </a:effectLst>'''))
        plain = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
        plain.text = "ふつうの文字"
        out = io.BytesIO()
        prs.save(out)
        deck = read_pptx_exact(out.getvalue())
        fx = by_text(deck["slideData"][0], "光る文字")
        self.assertEqual(fx["tglow"]["color"], "#b7c3da")
        self.assertEqual(fx["tglow"]["opacity"], 0.6)
        self.assertGreater(fx["tglow"]["r"], 5)
        self.assertEqual(fx["tshadow"]["color"], "#000000")
        self.assertEqual(fx["tshadow"]["opacity"], 0.4)
        self.assertGreater(fx["tshadow"]["dx"], 0)
        self.assertGreater(fx["tshadow"]["dy"], 0)
        self.assertGreater(fx["tshadow"]["blur"], 0)
        self.assertEqual(fx["toutline"]["color"], "#1f3864", "文字の輪郭")
        self.assertAlmostEqual(fx["toutline"]["w"], 3.0, places=1)
        self.assertNotIn("shadow", fx, "the shape's own shadow is another thing")
        ordinary = by_text(deck["slideData"][0], "ふつうの文字")
        self.assertNotIn("tshadow", ordinary)
        self.assertNotIn("tglow", ordinary)
        self.assertNotIn("toutline", ordinary)


if __name__ == "__main__":
    unittest.main()


class ReviewCommentsTest(unittest.TestCase):
    """PowerPoint's comments come over as threads with who wrote them, when, replies and the shape they are on."""

    @staticmethod
    def _part(prs, name, content_type, xml):
        from pptx.opc.package import Part
        from pptx.opc.packuri import PackURI
        return Part(PackURI(name), content_type, prs.part.package, blob=xml.encode("utf-8"))

    def test_modern_comments_keep_authors_replies_status_and_their_shape(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        box.text = "Target"
        p188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
        authors = self._part(prs, "/ppt/authors.xml", "application/vnd.ms-powerpoint.authors+xml",
                             f'<p188:authorLst xmlns:p188="{p188}"><p188:author id="{{A1}}" name="山田 太郎" initials="YT" userId="y" providerId="None"/>'
                             f'<p188:author id="{{B2}}" name="Sato" initials="S" userId="s" providerId="None"/></p188:authorLst>')
        prs.part.relate_to(authors, "http://schemas.microsoft.com/office/2018/10/relationships/authors")
        body = '<p188:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>{}</a:t></a:r></a:p></p188:txBody>'
        cm = self._part(prs, "/ppt/comments/modernComment_100_0.xml", "application/vnd.ms-powerpoint.comments+xml",
                        f'<p188:cmLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p188="{p188}">'
                        f'<p188:cm id="{{C1}}" authorId="{{A1}}" created="2026-09-01T10:00:00.000">'
                        '<ac:deMkLst xmlns:ac="http://schemas.microsoft.com/office/drawing/2013/main/command" xmlns:pc="http://schemas.microsoft.com/office/powerpoint/2013/main/command">'
                        f'<pc:docMk/><pc:sldMk cId="0" sldId="256"/><ac:spMk id="{box.shape_id}" creationId="{{X}}"/></ac:deMkLst>'
                        f'<p188:replyLst><p188:reply id="{{R1}}" authorId="{{B2}}" created="2026-09-02T09:00:00.000">{body.format("直しました")}</p188:reply></p188:replyLst>'
                        f'{body.format("数字を確認 @Sato")}</p188:cm>'
                        f'<p188:cm id="{{C2}}" authorId="{{B2}}" created="2026-09-03T09:00:00.000" status="resolved">{body.format("済み")}</p188:cm></p188:cmLst>')
        slide.part.relate_to(cm, "http://schemas.microsoft.com/office/2018/10/relationships/comments")
        out = io.BytesIO()
        prs.save(out)
        deck = read_pptx_exact(out.getvalue())
        page = deck["slideData"][0]
        first, second = page["comments"]
        self.assertEqual(first["text"], "数字を確認 @Sato")
        self.assertEqual(first["by"], "山田 太郎")
        self.assertTrue(first["at"].startswith("2026-09-01"))
        self.assertEqual(first["anchor"], by_text(page, "Target")["id"])
        self.assertEqual([(r["by"], r["text"]) for r in first["replies"]], [("Sato", "直しました")])
        self.assertTrue(second["done"])
        self.assertNotIn("done", first)

    def test_classic_comments_and_their_threaded_replies(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        authors = self._part(prs, "/ppt/commentAuthors.xml", "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml",
                             '<p:cmAuthorLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                             '<p:cmAuthor id="0" name="Reviewer A" initials="RA" lastIdx="1" clrIdx="0"/><p:cmAuthor id="1" name="Writer B" initials="WB" lastIdx="1" clrIdx="1"/></p:cmAuthorLst>')
        prs.part.relate_to(authors, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors")
        cm = self._part(prs, "/ppt/comments/comment1.xml", "application/vnd.openxmlformats-officedocument.presentationml.comments+xml",
                        '<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                        '<p:cm authorId="0" dt="2026-08-01T10:00:00.000" idx="1"><p:pos x="10" y="10"/><p:text>Shorter title?</p:text></p:cm>'
                        '<p:cm authorId="1" dt="2026-08-02T10:00:00.000" idx="1"><p:pos x="10" y="10"/><p:text>Done.</p:text>'
                        '<p:extLst><p:ext uri="{C676402C-5697-4E1C-873F-D02D1690AC5C}"><p15:threadingInfo xmlns:p15="http://schemas.microsoft.com/office/powerpoint/2012/main" timeZoneBias="-540">'
                        '<p15:parentCm authorId="0" idx="1"/></p15:threadingInfo></p:ext></p:extLst></p:cm></p:cmLst>')
        slide.part.relate_to(cm, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")
        out = io.BytesIO()
        prs.save(out)
        page = read_pptx_exact(out.getvalue())["slideData"][0]
        self.assertEqual(len(page["comments"]), 1)
        thread = page["comments"][0]
        self.assertEqual((thread["by"], thread["text"]), ("Reviewer A", "Shorter title?"))
        self.assertEqual([(r["by"], r["text"]) for r in thread["replies"]], [("Writer B", "Done.")])


class Model3dTest(unittest.TestCase):
    """A PowerPoint 3D model comes over as a 3D model: its GLB, its turn, the picture PowerPoint drew, its name."""

    def test_a_3d_model_keeps_its_file_turn_and_picture(self):
        from pptx.opc.package import Part
        from pptx.opc.packuri import PackURI
        import base64
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        glb = b"glTF" + (2).to_bytes(4, "little") + (12).to_bytes(4, "little")
        model = Part(PackURI("/ppt/media/model3d1.glb"), "model/gltf-binary", prs.part.package, blob=glb)
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
        picture = Part(PackURI("/ppt/media/image9.png"), "image/png", prs.part.package, blob=png)
        rid_model = slide.part.relate_to(model, "http://schemas.microsoft.com/office/2017/06/relationships/model3d")
        rid_pic = slide.part.relate_to(picture, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
        ns = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
              'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
              'xmlns:am3d="http://schemas.microsoft.com/office/drawing/2017/model3d"')
        xml = (f'<mc:AlternateContent {ns}><mc:Choice Requires="am3d"><p:graphicFrame>'
               '<p:nvGraphicFramePr><p:cNvPr id="7" name="3D モデル 6" descr="店舗の模型"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
               f'<p:xfrm><a:off x="{Inches(1)}" y="{Inches(1)}"/><a:ext cx="{Inches(3)}" cy="{Inches(2)}"/></p:xfrm>'
               '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2017/model3d">'
               f'<am3d:model3d r:embed="{rid_model}"><am3d:spPr/><am3d:camera/><am3d:trans><am3d:rot ax="1200000" ay="-1800000" az="0"/></am3d:trans>'
               f'<am3d:raster rName="Office3DRenderer" rVer="16.0.8326"><am3d:blip r:embed="{rid_pic}"/></am3d:raster></am3d:model3d>'
               '</a:graphicData></a:graphic></p:graphicFrame></mc:Choice>'
               f'<mc:Fallback><p:pic><p:nvPicPr><p:cNvPr id="7" name="3D モデル 6"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid_pic}"/></p:blipFill>'
               f'<p:spPr><a:xfrm><a:off x="{Inches(1)}" y="{Inches(1)}"/><a:ext cx="{Inches(3)}" cy="{Inches(2)}"/></a:xfrm></p:spPr></p:pic></mc:Fallback></mc:AlternateContent>')
        slide.shapes._spTree.append(etree.fromstring(xml))
        out = io.BytesIO()
        prs.save(out)
        deck = read_pptx_exact(out.getvalue())
        [o] = [x for x in deck["slideData"][0]["elements"] if x["kind"] == "model"]
        self.assertTrue(o["src"].startswith("data:model/gltf-binary;base64,"))
        self.assertEqual(base64.b64decode(o["src"].split(",", 1)[1]), glb)
        self.assertEqual(o["view"], {"pitch": 20.0, "yaw": -30.0})
        self.assertTrue(o["poster"].startswith("data:image/png"))
        self.assertEqual((o["name"], o["alt"]), ("3D モデル 6", "店舗の模型"))
        self.assertGreater(o["w"], o["h"], "its box keeps PowerPoint's shape")
        self.assertEqual(deck["stats"]["models"], 1)
        self.assertFalse(any(x["kind"] == "image" for x in deck["slideData"][0]["elements"]), "not also its fallback picture")


class ChartExTest(unittest.TestCase):
    """Office 2016's charts (chartEx) come over as the studio's own editable charts; a kind the studio does not draw
    keeps the picture PowerPoint stored with it."""

    CX = "http://schemas.microsoft.com/office/drawing/2014/chartex"

    def deck_with(self, charts):
        from pptx.opc.package import Part
        from pptx.opc.packuri import PackURI
        import base64
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
        picture = Part(PackURI("/ppt/media/image7.png"), "image/png", prs.part.package, blob=png)
        rid_pic = slide.part.relate_to(picture, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
        ns = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
              'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
              f'xmlns:cx="{self.CX}"')
        for k, body in enumerate(charts):
            part = Part(PackURI(f"/ppt/charts/chartEx{k + 1}.xml"), "application/vnd.ms-office.chartex+xml", prs.part.package,
                        blob=f'<cx:chartSpace xmlns:cx="{self.CX}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">{body}</cx:chartSpace>'.encode("utf-8"))
            rid = slide.part.relate_to(part, "http://schemas.microsoft.com/office/2014/relationships/chartEx")
            sid = 10 + k
            box = f'<a:off x="{Inches(1 + k)}" y="{Inches(1)}"/><a:ext cx="{Inches(4)}" cy="{Inches(3)}"/>'
            xml = (f'<mc:AlternateContent {ns}><mc:Choice Requires="cx1"><p:graphicFrame>'
                   f'<p:nvGraphicFramePr><p:cNvPr id="{sid}" name="グラフ {sid}"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
                   f'<p:xfrm>{box}</p:xfrm><a:graphic><a:graphicData uri="{self.CX}"><cx:chart r:id="{rid}"/></a:graphicData></a:graphic></p:graphicFrame></mc:Choice>'
                   f'<mc:Fallback><p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="グラフ {sid}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
                   f'<p:spPr><a:xfrm>{box}</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:blipFill><a:blip r:embed="{rid_pic}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill></p:spPr></p:sp>'
                   '</mc:Fallback></mc:AlternateContent>')
            slide.shapes._spTree.append(etree.fromstring(xml))
        out = io.BytesIO()
        prs.save(out)
        return read_pptx_exact(out.getvalue())

    @staticmethod
    def dim(kind, tag, levels):
        lv = "".join(f'<cx:lvl ptCount="{len(lvl)}">' + "".join(f'<cx:pt idx="{i}">{v}</cx:pt>' for i, v in enumerate(lvl) if v != "") + "</cx:lvl>" for lvl in levels)
        return f'<cx:{tag} type="{kind}"><cx:f>Sheet1!A1</cx:f>{lv}</cx:{tag}>'

    def chart(self, layout, data, series_extra="", title="", more_series=""):
        t = f'<cx:title><cx:tx><cx:txData><cx:v>{title}</cx:v></cx:txData></cx:tx></cx:title>' if title else ""
        return (f'<cx:chartData>{data}</cx:chartData><cx:chart>{t}<cx:plotArea><cx:plotAreaRegion>'
                f'<cx:series layoutId="{layout}" uniqueId="{{1}}"><cx:tx><cx:txData><cx:v>売上</cx:v></cx:txData></cx:tx><cx:dataId val="0"/>{series_extra}</cx:series>{more_series}'
                '</cx:plotAreaRegion></cx:plotArea></cx:chart>')

    def test_treemap_boxplot_waterfall_come_over_and_a_map_keeps_its_picture(self):
        treemap = self.chart("treemap", '<cx:data id="0">' + self.dim("cat", "strDim", [["おにぎり", "弁当", "お茶"], ["食品", "", "飲料"]]) + self.dim("size", "numDim", [["42", "35", "24"]]) + "</cx:data>", title="カテゴリ別売上")
        box = self.chart("boxWhisker", '<cx:data id="0">' + self.dim("cat", "strDim", [["東", "東", "東", "西", "西"]]) + self.dim("val", "numDim", [["60", "70", "80", "50", "90"]]) + "</cx:data>")
        fall = self.chart("waterfall", '<cx:data id="0">' + self.dim("cat", "strDim", [["前年", "増", "減", "今年"]]) + self.dim("val", "numDim", [["100", "30", "-10", "120"]]) + "</cx:data>",
                          series_extra='<cx:layoutPr><cx:subtotals><cx:idx val="0"/><cx:idx val="3"/></cx:subtotals></cx:layoutPr>')
        region = self.chart("regionMap", '<cx:data id="0">' + self.dim("cat", "strDim", [["東京", "大阪"]]) + self.dim("val", "numDim", [["5", "3"]]) + "</cx:data>")
        deck = self.deck_with([treemap, box, fall, region])
        els = deck["slideData"][0]["elements"]
        charts = [x["chart"] for x in els if x["kind"] == "chart"]
        self.assertEqual([c["type"] for c in charts], ["treemap", "boxplot", "waterfall"])
        tm, bx, wf = charts
        self.assertEqual(tm["labels"], ["食品/おにぎり", "食品/弁当", "飲料/お茶"], "the parent level carried down to its members")
        self.assertEqual(tm["series"][0]["values"], [42, 35, 24])
        self.assertEqual(tm["title"], "カテゴリ別売上")
        self.assertEqual([s["name"] for s in bx["series"]], ["東", "西"], "one series grouped by its categories")
        self.assertEqual(bx["series"][0]["values"], [60, 70, 80])
        self.assertEqual(bx["series"][1]["values"], [50, 90, None], "a shorter group padded with empty cells")
        self.assertEqual(wf["opts"], {"totals": [0, 3]})
        self.assertEqual(deck["stats"]["charts"], 3)
        self.assertEqual(len([x for x in els if x["kind"] != "chart"]), 1, "the map keeps PowerPoint's picture")


class DistributedAlignTest(unittest.TestCase):
    """均等割り付け (algn="dist") comes over as distributed, not as justified."""

    def test_a_distributed_paragraph_stays_distributed(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "均等割り付け"
        box.text_frame.add_paragraph().text = "左揃え"
        box.text_frame.paragraphs[0]._p.get_or_add_pPr().set("algn", "dist")
        out = io.BytesIO()
        prs.save(out)
        deck = read_pptx_exact(out.getvalue())
        imported = by_text(deck["slideData"][0], "均等割り付け")
        self.assertEqual(imported["align"], "distributed")
        self.assertIn('style="text-align: justify; text-align-last: justify"', imported["text"])


class MediaTest(unittest.TestCase):
    """PowerPoint's videos come over playable, with their poster, trim, fades, bookmarks and volume."""

    @staticmethod
    def mp4(seconds):
        import struct
        ftyp = struct.pack(">I4s4sI4s", 20, b"ftyp", b"isom", 0, b"isom")
        body = bytes(4) + struct.pack(">IIII", 0, 0, 1000, int(seconds * 1000)) + bytes(80)
        mvhd = struct.pack(">I4s", 8 + len(body), b"mvhd") + body
        return ftyp + struct.pack(">I4s", 8 + len(mvhd), b"moov") + mvhd

    def deck_with(self, blob, mime, suffix=".mp4"):
        import tempfile
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with tempfile.NamedTemporaryFile(suffix=suffix) as f:
            f.write(blob)
            f.flush()
            pic = slide.shapes.add_movie(f.name, Inches(1), Inches(1), Inches(4), Inches(2.25), mime_type=mime)
        p14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
        media = next(pic._element.iter(f"{{{p14}}}media"), None)
        if media is not None:
            for xml in (f'<p14:trim xmlns:p14="{p14}" st="1000" end="1500"/>', f'<p14:fade xmlns:p14="{p14}" in="500" out="250"/>',
                        f'<p14:bmkLst xmlns:p14="{p14}"><p14:bmk name="サビ" time="2500"/><p14:bmk name="終わり" time="4000"/></p14:bmkLst>'):
                media.append(etree.fromstring(xml))
        out = io.BytesIO()
        prs.save(out)
        return read_pptx_exact(out.getvalue())

    def test_an_mp4_video_plays_with_its_settings(self):
        deck = self.deck_with(self.mp4(6), "video/mp4")
        video = next(o for o in deck["slideData"][0]["elements"] if o["kind"] == "video")
        self.assertTrue(video["src"].startswith("data:video/mp4;base64,"))
        self.assertTrue(video["fileName"].endswith(".mp4"))
        self.assertEqual((video["autoplay"], video["loop"], video["muted"]), (False, False, False), "plays on a click, with sound")
        self.assertEqual(video["trimStart"], 1.0)
        self.assertEqual(video["trimEnd"], 4.5, "6 s long, 1.5 s cut from the end")
        self.assertEqual((video["fadeIn"], video["fadeOut"]), (0.5, 0.25))
        self.assertEqual(video["bookmarks"], [{"name": "サビ", "t": 2.5}, {"name": "終わり", "t": 4.0}])
        self.assertEqual(video.get("volume"), 0.8, "python-pptx's timing keeps PowerPoint's 80% volume")
        self.assertTrue(video.get("poster", "").startswith("data:image/"), "the poster frame")
        self.assertEqual(deck["stats"]["media"], 1)

    def test_a_video_browsers_cannot_play_stays_its_picture(self):
        deck = self.deck_with(b"0&\xb2u\x8ef\xcf\x11" + bytes(64), "video/x-ms-wmv", ".wmv")
        kinds = [o["kind"] for o in deck["slideData"][0]["elements"]]
        self.assertNotIn("video", kinds)
        self.assertIn("image", kinds, "the poster stays on the slide")

    def test_media_lengths_from_their_headers(self):
        import struct
        self.assertAlmostEqual(pptx_exact.media_duration(self.mp4(12.5), ".mp4"), 12.5)
        rate = 8000 * 2
        wav = b"RIFF" + struct.pack("<I", 36 + rate * 3) + b"WAVE" + b"fmt " + struct.pack("<IHHIIHH", 16, 1, 1, 8000, rate, 2, 16) + b"data" + struct.pack("<I", rate * 3) + bytes(rate * 3)
        self.assertAlmostEqual(pptx_exact.media_duration(wav, ".wav"), 3.0)
        self.assertIsNone(pptx_exact.media_duration(b"nothing", ".mp3"))
