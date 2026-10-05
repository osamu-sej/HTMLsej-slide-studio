"""PowerPoint decks for the faithful import (tools/pptx_exact.py): one in the SEJ template and one plain deck.

They carry what an imported deck must keep: shapes with text runs, bullets and a link, a table, charts with their
own colours and data labels, a picture, a connector with an arrowhead, a group, rotation, speaker notes, an
entrance animation, a slide transition and a hidden slide.
Usage (for the browser check): python test/pptx_fixtures.py <folder>   → sej.pptx, plain43.pptx
"""
import io
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
NAVY, PALE, GREY_BLUE, BROWN = RGBColor(0x1F, 0x38, 0x64), RGBColor(0xDC, 0xE4, 0xF2), RGBColor(0xB7, 0xC3, 0xDA), RGBColor(0xF5, 0xF0, 0xEA)


def photo() -> io.BytesIO:
    img = Image.new("RGB", (800, 500), (220, 228, 242))
    draw = ImageDraw.Draw(img)
    for i in range(0, 800, 40):
        draw.rectangle([i, 300 - i // 4, i + 30, 500], fill=(31, 56, 100))
    draw.ellipse([560, 40, 740, 220], fill=(214, 201, 184))
    out = io.BytesIO()
    img.save(out, "PNG")
    out.seek(0)
    return out


def animate(slide, effects, transition="fade"):
    """A slide transition and entrance effects [(shape, presetID, presetSubtype, nodeType)] in PowerPoint's timing tree."""
    sld = slide._element
    if transition:
        tr = etree.SubElement(sld, f"{{{P_NS}}}transition", spd="med")
        etree.SubElement(tr, f"{{{P_NS}}}{transition}")
    ids = iter(range(3, 999))
    pars = []
    for shape, preset, sub, node in effects:
        spid = str(shape.shape_id)
        pars.append(f"""<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="{'indefinite' if node == 'clickEffect' else '0'}"/></p:stCondLst><p:childTnLst>
          <p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
            <p:par><p:cTn id="{next(ids)}" presetID="{preset}" presetClass="entr" presetSubtype="{sub}" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
              <p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
              <p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{next(ids)}" dur="500"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>
            </p:childTnLst></p:cTn></p:par>
          </p:childTnLst></p:cTn></p:par>
        </p:childTnLst></p:cTn></p:par>""")
    xml = f"""<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
      <p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{''.join(pars)}</p:childTnLst></p:cTn>
      <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
      <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
    </p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""
    sld.append(etree.fromstring(xml))


def run(paragraph, text, size, bold=False, color=None):
    r = paragraph.add_run()
    r.text, r.font.size, r.font.bold = text, Pt(size), bold
    if color is not None:
        r.font.color.rgb = color
    return r


def content(slide, sej: bool):
    """The shapes every deck carries; returns the ones the animation targets."""
    shapes = slide.shapes
    box = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(3.6), Inches(1.6))
    box.name = "課題"
    box.fill.solid()
    box.fill.fore_color.rgb = PALE
    box.line.fill.background()
    tf = box.text_frame
    run(tf.paragraphs[0], "現状の課題", 20, True, NAVY)
    run(tf.add_paragraph(), "発注に1日2時間かかる", 14, False, RGBColor(0x1A, 0x1A, 0x1A))
    arrow = shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.4), Inches(2.1), Inches(1.0), Inches(0.6))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = GREY_BLUE
    arrow.line.fill.background()
    goal = shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.6), Inches(1.5), Inches(1.8), Inches(1.8))
    goal.fill.solid()
    goal.fill.fore_color.rgb = BROWN
    goal.line.fill.background()
    goal.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    run(goal.text_frame.paragraphs[0], "30分", 28, True, NAVY)
    tb = shapes.add_textbox(Inches(0.6), Inches(3.5), Inches(6.8), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(["端末で発注する（写真で確認）", "検品はバーコードで行う", "例外だけを人が見る"]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run(para, line, 16)
        ppr = para._p.get_or_add_pPr()
        ppr.set("marL", "228600")
        ppr.set("indent", "-228600")
        ppr.append(ppr.makeelement(f"{A}buChar", {"char": "■"}))
    para = tf.add_paragraph()
    run(para, "詳しくは ", 14)
    run(para, "社内ポータル", 14).hyperlink.address = "https://example.com/portal"
    table = shapes.add_table(3, 3, Inches(7.8), Inches(1.5), Inches(5.0), Inches(1.6)).table
    for r, row in enumerate([["項目", "現状", "目標"], ["作業時間", "120h", "80h"], ["廃棄率", "3.0%", "2.5%"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    data = CategoryChartData()
    data.categories = ["A店", "B店", "C店"]
    data.add_series("削減時間", (42, 30, 18))
    chart = shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(7.8), Inches(3.4), Inches(5.0), Inches(2.6), data).chart
    chart.font.size = Pt(14)
    plot = chart.plots[0]
    plot.gap_width = 60
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = GREY_BLUE
    series.points[0].format.fill.solid()
    series.points[0].format.fill.fore_color.rgb = NAVY
    plot.has_data_labels = True
    plot.data_labels.number_format = '0"h"'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    chart.value_axis.visible = False
    if not sej:
        shapes.add_picture(photo(), Inches(10.9), Inches(0.2), Inches(1.6), Inches(1.0))
    line = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6), Inches(6.0), Inches(7.4), Inches(6.0))
    line.line.color.rgb = NAVY
    line.line.width = Pt(2)
    ln = line.line._get_or_add_ln()
    ln.append(ln.makeelement(f"{A}tailEnd", {"type": "triangle"}))
    group = shapes.add_group_shape()
    for i, label in enumerate(["グループ1", "グループ2"]):
        member = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6 + i * 1.4), Inches(6.2), Inches(1.2), Inches(0.5))
        member.fill.solid()
        member.fill.fore_color.rgb = PALE
        member.line.fill.background()
        member.text_frame.text = label
    return box, goal


def save(prs) -> bytes:
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def sej_deck() -> bytes:
    prs = Presentation(str(ROOT / "assets" / "sej" / "template.pptx"))
    ids = prs.slides._sldIdLst
    for sld_id in list(ids):  # the template's own sample slides
        prs.part.drop_rel(sld_id.rId)
        ids.remove(sld_id)
    cover = prs.slides.add_slide(prs.slide_masters[0].slide_layouts[0])
    for ph in cover.placeholders:
        ph.text = {1: "発注業務の改革", 4: "オペレーション情報部"}.get(ph.placeholder_format.type, "2026年10月")
    layout = next(l for l in prs.slide_masters[1].slide_layouts if l.name == "タイトルとコンテンツ")
    body = prs.slides.add_slide(layout)
    for ph in list(body.placeholders):
        if ph.placeholder_format.type == 1:
            ph.text = "改革の全体像"
        else:
            ph._element.getparent().remove(ph._element)
    box, goal = content(body, True)
    body.notes_slide.notes_text_frame.text = "ここで全体像を説明します。"
    animate(body, [(box, 10, 0, "clickEffect"), (goal, 2, 8, "afterEffect")])
    hidden = prs.slides.add_slide(layout)
    hidden._element.set("show", "0")
    hidden.shapes.title.text = "隠したページ"
    points = prs.slides.add_slide(layout)
    for ph in points.placeholders:
        if ph.placeholder_format.type == 1:
            ph.text = "箇条書き"
        else:
            ph.text_frame.text = "最初の点"
            for words, level in (("詳しい点", 1), ("もう一つの点", 0)):
                para = ph.text_frame.add_paragraph()
                para.text, para.level = words, level
    # Buttons a slide show follows: back to the overview page, and on to the end.
    back = points.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.5), Inches(5.6), Inches(2.6), Inches(0.6))
    back.text_frame.text = "全体像へ戻る"
    back.click_action.target_slide = body
    # 線の書式設定: a thick double line with round ends and bevelled corners.
    back.line.width = Pt(6)
    back.line.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    ln = back._element.spPr.find(f"{A}ln")
    ln.set("cap", "rnd")
    ln.set("cmpd", "dbl")
    ln.append(ln.makeelement(f"{A}bevel", {}))
    # マウスの通過: moving the pointer onto it goes on to the next slide.
    bcnv = back._element.nvSpPr.cNvPr
    bcnv.append(bcnv.makeelement(f"{A}hlinkHover", {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": "", "action": "ppaction://hlinkshowjump?jump=nextslide"}))
    end = points.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(5.6), Inches(2.6), Inches(0.6))
    end.text_frame.text = "終わる"
    cnv = end._element.nvSpPr.cNvPr
    cnv.append(cnv.makeelement(f"{A}hlinkClick", {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": "", "action": "ppaction://hlinkshowjump?jump=endshow"}))
    return save(prs)


def plain_deck() -> bytes:
    """A 4:3 deck in PowerPoint's default template (pillarboxed on the studio's 16:9 page)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "四半期の報告"
    cover.placeholders[1].text = "営業企画部"
    fill = cover.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xFB)
    body = prs.slides.add_slide(prs.slide_layouts[5])
    body.shapes.title.text = "全体像"
    content(body, False)
    shapes = prs.slides.add_slide(prs.slide_layouts[6]).shapes
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("構成", (50, 30, 20))
    pie = shapes.add_chart(XL_CHART_TYPE.PIE, Inches(1), Inches(1), Inches(4), Inches(4), data).chart
    pie.plots[0].has_data_labels = True
    pie.plots[0].data_labels.show_percentage = True
    pie.plots[0].data_labels.show_value = False
    star = shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(6), Inches(1), Inches(2), Inches(2))
    chevron = shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(6), Inches(3.5), Inches(2.5), Inches(1))
    chevron.rotation = 20
    shapes.add_picture(photo(), Inches(1), Inches(5.2), Inches(3), Inches(1.8))
    assert star
    return save(prs)


if __name__ == "__main__":
    folder = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "sej.pptx").write_bytes(sej_deck())
    (folder / "plain43.pptx").write_bytes(plain_deck())
    print(folder / "sej.pptx", folder / "plain43.pptx")
