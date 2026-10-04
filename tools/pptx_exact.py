"""Bring a PowerPoint deck over as it looks (見た目どおりに取り込む).

Every slide becomes a 白紙 page of objects the studio draws and edits (public/engine/objects.js): shapes with
their outline, fill, outline colour and width, rotation and flips; the text inside them run by run (size, bold,
italic, underline, colour, links), paragraph by paragraph (alignment, line spacing, bullets and numbers); text
boxes; pictures with their crop; lines and connectors with arrowheads; tables cell by cell (fills, borders,
merged cells); charts with their data; SmartArt from its drawing; groups; the background. Positions keep the
slide's proportions on the studio's 1920 × 1080 page. Placeholders take what their layout and master give
them (position, size, colour, bullets), colours follow the theme. Slide transitions and entrance / exit /
emphasis animations come along as the studio's own.

The source template's marks, backgrounds and decorations remain editable objects on every imported page.
The studio's own master is not drawn over them.
"""

from __future__ import annotations

import base64
import colorsys
import io
import math
import posixpath
import re
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from lxml import etree
from PIL import Image

W, H = 1920, 1080
EMU_PT = 12700
NS = {
    "cx": "http://schemas.microsoft.com/office/drawing/2014/chartex",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "dsp": "http://schemas.microsoft.com/office/drawing/2008/diagram",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "am3d": "http://schemas.microsoft.com/office/drawing/2017/model3d",
}
A = "{%s}" % NS["a"]
AM3D = "{%s}" % NS["am3d"]
P = "{%s}" % NS["p"]
R = "{%s}" % NS["r"]
DSP = "{%s}" % NS["dsp"]

# PowerPoint's preset shapes → the studio's (public/engine/objects.js SHAPES). `scale` turns an OOXML
# adjustment (1/100000) into ours; "deg" angles (1/60000 degree).
SHAPE_MAP = {name: name for name in (
    "rect roundRect snip1Rect snip2SameRect snip2DiagRect round1Rect round2SameRect round2DiagRect ellipse triangle "
    "rtTriangle parallelogram trapezoid diamond pentagon hexagon heptagon octagon decagon dodecagon pie chord teardrop "
    "frame halfFrame corner diagStripe plus plaque can cube bevel donut noSmoking blockArc foldedCorner smileyFace heart "
    "lightningBolt sun moon cloud arc bracketPair bracePair leftBracket rightBracket leftBrace rightBrace rightArrow "
    "leftArrow upArrow downArrow leftRightArrow upDownArrow quadArrow bentArrow uturnArrow bentUpArrow circularArrow "
    "stripedRightArrow notchedRightArrow homePlate chevron rightArrowCallout downArrowCallout mathPlus mathMinus "
    "mathMultiply mathDivide mathEqual mathNotEqual star4 star5 star6 star7 star8 star10 star12 star16 star24 star32 "
    "wave doubleWave wedgeRectCallout wedgeRoundRectCallout wedgeEllipseCallout cloudCallout").split()}
SHAPE_MAP.update({
    "flowChartProcess": "flowProcess", "flowChartAlternateProcess": "flowAlternate", "flowChartDecision": "flowDecision",
    "flowChartInputOutput": "flowData", "flowChartPredefinedProcess": "flowPredefined", "flowChartInternalStorage": "flowInternalStorage",
    "flowChartDocument": "flowDocument", "flowChartMultidocument": "flowMultidocument", "flowChartTerminator": "flowTerminator",
    "flowChartPreparation": "flowPreparation", "flowChartManualInput": "flowManualInput", "flowChartManualOperation": "flowManualOperation",
    "flowChartConnector": "flowConnector", "flowChartOffpageConnector": "flowOffpage", "flowChartPunchedCard": "flowCard",
    "flowChartPunchedTape": "flowPunchedTape", "flowChartSummingJunction": "flowSummingJunction", "flowChartOr": "flowOr",
    "flowChartCollate": "flowCollate", "flowChartSort": "flowSort", "flowChartExtract": "flowExtract", "flowChartMerge": "flowMerge",
    "flowChartOnlineStorage": "flowStoredData", "flowChartDelay": "flowDelay", "flowChartMagneticDisk": "flowMagneticDisk",
    "flowChartMagneticDrum": "flowDirectAccess", "flowChartDisplay": "flowDisplay", "flowChartMagneticTape": "flowProcess",
    "flowChartOfflineStorage": "flowMerge", "irregularSeal1": "explosion1", "irregularSeal2": "explosion2",
    "snipRoundRect": "round1Rect", "nonIsoscelesTrapezoid": "trapezoid", "leftCircularArrow": "circularArrow",
    "borderCallout1": "lineCallout", "borderCallout2": "lineCallout", "borderCallout3": "lineCallout", "callout1": "lineCallout",
    "callout2": "lineCallout", "callout3": "lineCallout", "accentCallout1": "lineCallout", "accentBorderCallout1": "lineCallout",
    "leftArrowCallout": "rightArrowCallout", "rightArrowCallout": "rightArrowCallout", "upArrowCallout": "downArrowCallout",
    "leftRightArrowCallout": "rightArrowCallout", "upDownArrowCallout": "downArrowCallout", "quadArrowCallout": "quadArrow",
    "curvedRightArrow": "rightArrow", "curvedLeftArrow": "leftArrow", "curvedUpArrow": "upArrow", "curvedDownArrow": "downArrow",
    "leftUpArrow": "bentUpArrow", "leftRightUpArrow": "quadArrow", "swooshArrow": "rightArrow", "ellipseRibbon": "wave",
    "ellipseRibbon2": "wave", "ribbon": "rect", "ribbon2": "rect", "verticalScroll": "foldedCorner", "horizontalScroll": "foldedCorner",
    "doubleWave": "doubleWave", "pieWedge": "pie", "funnel": "trapezoid", "gear6": "sun", "gear9": "sun", "squareTabs": "rect",
    "cornerTabs": "rect", "plaqueTabs": "rect", "chartPlus": "plus", "chartX": "mathMultiply", "chartStar": "star6",
    "actionButtonBlank": "rect", "flowChartPunchedCard ": "flowCard",
})
# Adjustments that mean the same in both, as a factor; None drops them (the studio's default is used).
ADJ_SCALE = {key: 1.0 for key in (
    "roundRect snip1Rect snip2SameRect snip2DiagRect round1Rect round2SameRect round2DiagRect triangle parallelogram trapezoid "
    "hexagon octagon teardrop frame halfFrame corner diagStripe plus plaque can cube bevel donut noSmoking foldedCorner "
    "smileyFace sun moon bracketPair bracePair leftBracket rightBracket leftBrace rightBrace rightArrow leftArrow upArrow "
    "downArrow leftRightArrow upDownArrow quadArrow bentArrow uturnArrow bentUpArrow stripedRightArrow notchedRightArrow "
    "homePlate chevron mathPlus mathMinus mathMultiply mathDivide mathEqual mathNotEqual wave doubleWave wedgeRectCallout "
    "wedgeRoundRectCallout wedgeEllipseCallout cloudCallout").split()}
ADJ_SCALE.update({f"star{n}": 2.0 for n in (4, 5, 6, 7, 8, 10, 12, 16, 24, 32)})
ADJ_ANGLES = {"pie": [True, True], "chord": [True, True], "arc": [True, True], "blockArc": [True, True, False]}
ADJ_COUNT = {"bentArrow": 3, "uturnArrow": 3, "bentUpArrow": 3, "quadArrow": 3, "leftBrace": 1, "rightBrace": 1, "wedgeRoundRectCallout": 2, "wave": 1, "doubleWave": 1}

DASHES = {"solid": None, "dot": "roundDot", "sysDot": "squareDot", "dash": "dash", "sysDash": "dash", "lgDash": "longDash",
          "dashDot": "dashDot", "sysDashDot": "dashDot", "lgDashDot": "longDashDot", "lgDashDotDot": "longDashDotDot", "sysDashDotDot": "longDashDotDot"}
ARROWS = {"triangle": "triangle", "arrow": "arrow", "stealth": "stealth", "diamond": "diamond", "oval": "oval"}
NAMED = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF", "yellow": "FFFF00",
         "gray": "808080", "grey": "808080", "silver": "C0C0C0", "navy": "000080", "darkBlue": "00008B", "orange": "FFA500"}
SEJ_TEXT = re.compile(r"(明日の笑顔|社内限り|秘（?[A-CＡ-Ｃ]）?|SEVEN-ELEVEN JAPAN|セブン‐イレブン・ジャパン)", re.I)
# A font's single line spacing (PowerPoint: ascent + descent), so text keeps its line positions.
# Bullets PowerPoint draws from Wingdings / Symbol, as the characters they look like.
WINGDINGS = {"§": "■", "Ø": "➢", "ü": "✓", "n": "■", "q": "❑", "l": "●", "p": "□", "v": "❖", "Ü": "➤", "ð": "⇒", "o": "□", "Ÿ": "•", "è": "➔", "\uf0a7": "■", "\uf0d8": "➢", "\uf0fc": "✓", "\uf06e": "■", "\uf0b7": "•"}
LINE_FACTORS = [(re.compile(r"Meiryo UI|メイリオ UI"), 1.33), (re.compile(r"Meiryo|メイリオ"), 1.5), (re.compile(r"游ゴシック|Yu Gothic|游明朝|Yu Mincho"), 1.44),
                (re.compile(r"ＭＳ|MS (P)?Gothic|MS (P)?Mincho|MS UI Gothic"), 1.0), (re.compile(r"BIZ UD"), 1.33), (re.compile(r"Noto Sans (JP|CJK)"), 1.45)]


def r2(v: float) -> float:
    return round(float(v), 2)


def clamp(v: float, lo: float, hi: float) -> float:
    return max(lo, min(hi, v))


def esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def local(el) -> str:
    return etree.QName(el).localname


# ---------------------------------------------------------------- colours

def _hsl_mod(rgb, lum_mod=1.0, lum_off=0.0, sat_mod=1.0):
    r, g, b = (c / 255 for c in rgb)
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = clamp(l * lum_mod + lum_off, 0, 1)
    s = clamp(s * sat_mod, 0, 1)
    return tuple(round(c * 255) for c in colorsys.hls_to_rgb(h, l, s))


class Theme:
    """A slide master's theme: colours (with the master's colour map), fonts and the shape styles."""

    def __init__(self, master_part, master_el):
        self.colors: dict[str, str] = {}
        self.fonts = {"major": {}, "minor": {}}
        self.fills: list = []
        self.lines: list = []
        self.effects: list = []
        self.bg_fills: list = []
        theme_el = None
        for rel in master_part.rels.values():
            if rel.reltype.endswith("/theme"):
                theme_el = etree.fromstring(rel.target_part.blob)
        if theme_el is not None:
            scheme = theme_el.find(f".//{A}clrScheme")
            for child in scheme if scheme is not None else []:
                c = child[0] if len(child) else None
                if c is None:
                    continue
                self.colors[local(child)] = c.get("lastClr") if local(c) == "sysClr" else c.get("val", "000000")
            for kind in ("major", "minor"):
                font = theme_el.find(f".//{A}fontScheme/{A}{kind}Font")
                if font is not None:
                    for tag in ("latin", "ea"):
                        el = font.find(A + tag)
                        if el is not None and el.get("typeface"):
                            self.fonts[kind][tag] = el.get("typeface")
                    for el in font.findall(A + "font"):
                        if el.get("script") == "Jpan":
                            self.fonts[kind]["ea"] = el.get("typeface")
            fmt = theme_el.find(f".//{A}fmtScheme")
            if fmt is not None:
                self.fills = list(fmt.find(A + "fillStyleLst") if fmt.find(A + "fillStyleLst") is not None else [])
                self.lines = list(fmt.find(A + "lnStyleLst") if fmt.find(A + "lnStyleLst") is not None else [])
                self.effects = list(fmt.find(A + "effectStyleLst") if fmt.find(A + "effectStyleLst") is not None else [])
                self.bg_fills = list(fmt.find(A + "bgFillStyleLst") if fmt.find(A + "bgFillStyleLst") is not None else [])
        clr_map = master_el.find(P + "clrMap")
        self.map = dict(clr_map.attrib) if clr_map is not None else {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}

    def font(self, typeface: str | None) -> str | None:
        if not typeface:
            return None
        m = re.match(r"^\+(mj|mn)-(lt|ea|cs)$", typeface)
        if m:
            fonts = self.fonts["major" if m.group(1) == "mj" else "minor"]
            return fonts.get("ea" if m.group(2) == "ea" else "latin")
        return typeface


def color_of(container, theme: Theme, ph: tuple | None = None, clr_map: dict | None = None):
    """The colour a fill / text colour element gives: ((r, g, b), alpha), or None."""
    if container is None:
        return None
    el = None
    for child in container:
        if local(child) in ("srgbClr", "schemeClr", "sysClr", "prstClr", "scrgbClr", "hslClr"):
            el = child
            break
    if el is None:
        return None
    kind = local(el)
    rgb = None
    if kind == "srgbClr":
        rgb = el.get("val", "000000")
    elif kind == "sysClr":
        rgb = el.get("lastClr") or ("FFFFFF" if el.get("val") == "window" else "000000")
    elif kind == "prstClr":
        rgb = NAMED.get(el.get("val"), "000000")
    elif kind == "scrgbClr":
        rgb = "".join(f"{round(int(el.get(k, 0)) / 100000 * 255):02X}" for k in ("r", "g", "b"))
    elif kind == "hslClr":
        h, s, l = int(el.get("hue", 0)) / 21600000, int(el.get("sat", 0)) / 100000, int(el.get("lum", 0)) / 100000
        rgb = "".join(f"{round(c * 255):02X}" for c in colorsys.hls_to_rgb(h, l, s))
    elif kind == "schemeClr":
        val = el.get("val")
        if val == "phClr":
            if ph is None:
                return None
            base, base_alpha = ph
            rgb = "".join(f"{c:02X}" for c in base)
        else:
            mapping = clr_map or theme.map
            val = mapping.get(val, val)
            rgb = theme.colors.get(val, "000000")
    try:
        rgb = tuple(int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    except (TypeError, ValueError):
        return None
    alpha = 1.0
    lum_mod, lum_off, sat_mod = 1.0, 0.0, 1.0
    for mod in el:
        name, val = local(mod), int(mod.get("val", 0))
        if name == "alpha":
            alpha = val / 100000
        elif name == "lumMod":
            lum_mod = val / 100000
        elif name == "lumOff":
            lum_off = val / 100000
        elif name == "satMod":
            sat_mod = val / 100000
        elif name == "tint":
            rgb = tuple(round(c + (255 - c) * (1 - val / 100000)) for c in rgb)
        elif name == "shade":
            rgb = tuple(round(c * val / 100000) for c in rgb)
    if lum_mod != 1.0 or lum_off or sat_mod != 1.0:
        rgb = _hsl_mod(rgb, lum_mod, lum_off, sat_mod)
    return rgb, alpha


def hexc(color) -> str:
    return "#" + "".join(f"{clamp(c, 0, 255):02x}" for c in color[0])


# ---------------------------------------------------------------- the deck

class Deck:
    def __init__(self, prs, part_by_name):
        self.prs = prs
        self.cx, self.cy = prs.slide_width or 12192000, prs.slide_height or 6858000
        self.k = min(W / self.cx, H / self.cy)
        self.ox, self.oy = (W - self.cx * self.k) / 2, (H - self.cy * self.k) / 2
        self.part_by_name = part_by_name
        pres_el = prs.part._element
        self.default_text = pres_el.find(P + "defaultTextStyle")
        self.table_styles = {}
        for rel in prs.part.rels.values():
            if rel.reltype.endswith("/tableStyles"):
                root = etree.fromstring(rel.target_part.blob)
                for style in root.findall(A + "tblStyle"):
                    self.table_styles[style.get("styleId")] = style
        # PowerPoint's default table style, for a deck that names it without carrying its definition.
        if DEFAULT_TABLE_STYLE not in self.table_styles:
            self.table_styles[DEFAULT_TABLE_STYLE] = etree.fromstring(MEDIUM_STYLE_2)
        self.images: dict[str, str | None] = {}
        self.themes: dict[str, Theme] = {}
        self.sids: dict[str, str] = {}
        self.stats = {"slides": 0, "objects": 0, "pictures": 0, "tables": 0, "charts": 0, "animations": 0, "skipped": 0, "hidden": 0, "unsupported": 0, "comments": 0, "models": 0}
        self.sej = self._is_sej()
        self.comment_authors = comment_authors(prs)

    def _is_sej(self) -> bool:
        for master in self.prs.slide_masters:
            words = " ".join(sh.text_frame.text for sh in master.shapes if sh.has_text_frame)
            if SEJ_TEXT.search(words) and "社内限り" in words:
                return True
        return False

    def px(self, emu) -> float:
        return float(emu) * self.k

    def x(self, emu) -> float:
        return self.ox + float(emu) * self.k

    def y(self, emu) -> float:
        return self.oy + float(emu) * self.k

    def pt(self, points: float) -> float:
        return points * EMU_PT * self.k


def image_url(part, deck: Deck) -> str | None:
    """Keep web-native picture bytes unchanged, avoiding JPEG and resize losses."""
    key = part.partname
    if key in deck.images:
        return deck.images[key]
    url = None
    blob = part.blob
    ext = posixpath.splitext(str(key))[1].lower()
    native = {".svg": "image/svg+xml", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
              ".gif": "image/gif", ".webp": "image/webp", ".bmp": "image/bmp"}
    if ext in native:
        url = f"data:{native[ext]};base64," + base64.b64encode(blob).decode()
    else:
        try:
            with Image.open(io.BytesIO(blob)) as img:
                img.load()
                if img.format in ("WMF", "EMF"):
                    raise ValueError("vector metafile")
                alpha = img.mode in ("RGBA", "LA", "P") and (img.mode != "P" or "transparency" in img.info)
                img = img.convert("RGBA" if alpha else "RGB")
                out = io.BytesIO()
                if alpha:
                    img.save(out, "PNG", optimize=True)
                    url = "data:image/png;base64," + base64.b64encode(out.getvalue()).decode()
                else:
                    img.save(out, "PNG", optimize=True)
                    url = "data:image/png;base64," + base64.b64encode(out.getvalue()).decode()
        except Exception:
            url = None
    deck.images[key] = url
    return url


# ---------------------------------------------------------------- a slide

class SlideReader:
    def __init__(self, deck: Deck, slide, index: int, number: int):
        self.deck = deck
        self.slide = slide
        self.index = index
        self.number = number
        self.layout = slide.slide_layout
        self.master = self.layout.slide_master
        key = str(self.master.part.partname)
        if key not in deck.themes:
            deck.themes[key] = Theme(self.master.part, self.master._element)
        self.theme = deck.themes[key]
        self.clr_map = dict(self.theme.map)
        override = slide._element.find(f"{P}clrMapOvr/{A}overrideClrMapping")
        if override is not None:
            self.clr_map = dict(override.attrib)
        self.objects: list[dict[str, Any]] = []
        self.by_spid: dict[str, str] = {}
        self.actions: dict[str, dict[str, str]] = {}
        self.used: set[str] = set()
        self.title = ""
        self.importing_decoration = False

    # ---- helpers

    def color(self, container, ph=None):
        return color_of(container, self.theme, ph, self.clr_map)

    def new_id(self, part: str, spid: str | None) -> str:
        base = re.sub(r"[^A-Za-z0-9]", "", f"s{self.index}{part}{spid or ''}")[:28] or f"s{self.index}"
        out, n = base, 1
        while out in self.used:
            n += 1
            out = f"{base[:26]}{n}"
        self.used.add(out)
        return out

    def rel_target(self, part, rid: str):
        try:
            return part.related_part(rid)
        except (KeyError, AttributeError):
            return None

    # ---- placeholders: what the layout and master give

    def ph_info(self, el):
        ph = el.find(f".//{P}nvPr/{P}ph")
        if ph is None:
            return None
        return {"type": ph.get("type", "body" if ph.get("idx") else "obj"), "idx": ph.get("idx")}

    def inherited_ph(self, info):
        """The layout and master placeholders a slide placeholder takes its settings from."""
        if not info:
            return []
        out = []

        def find(shapes_el, match_idx=True):
            for sp in shapes_el.iter(f"{P}sp"):
                other = self.ph_info(sp)
                if other is None:
                    continue
                if match_idx and info.get("idx") is not None and other.get("idx") == info["idx"]:
                    return sp
            want = {"ctrTitle": "title", "subTitle": "body", "obj": "body"}.get(info["type"], info["type"])
            for sp in shapes_el.iter(f"{P}sp"):
                other = self.ph_info(sp)
                if other is None:
                    continue
                kind = {"ctrTitle": "title", "subTitle": "body", "obj": "body"}.get(other["type"], other["type"])
                if kind == want:
                    return sp
            return None

        layout_sp = find(self.layout._element)
        if layout_sp is not None:
            out.append(layout_sp)
        master_sp = find(self.master._element, match_idx=False)
        if master_sp is not None:
            out.append(master_sp)
        return out

    # ---- geometry

    def xfrm_of(self, el, inherited=()):
        for node in (el, *inherited):
            sppr = node.find(P + "spPr") if node.find(P + "spPr") is not None else node.find(DSP + "spPr")
            xfrm = sppr.find(A + "xfrm") if sppr is not None else None
            if xfrm is None:
                xfrm = node.find(P + "xfrm")  # graphic frames
            if xfrm is not None and xfrm.find(A + "off") is not None:
                off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
                return {"x": int(off.get("x", 0)), "y": int(off.get("y", 0)), "w": int(ext.get("cx", 0)) if ext is not None else 0,
                        "h": int(ext.get("cy", 0)) if ext is not None else 0, "rot": int(xfrm.get("rot", 0)) / 60000,
                        "flipH": xfrm.get("flipH") in ("1", "true"), "flipV": xfrm.get("flipV") in ("1", "true")}
        return None

    def box(self, xf, tf):
        """A shape's box (EMU) through its groups' transforms → slide pixels, rotation and flips."""
        x, y, w, h, rot, fh, fv = xf["x"], xf["y"], xf["w"], xf["h"], xf["rot"], xf["flipH"], xf["flipV"]
        for g in tf:
            cx, cy = x + w / 2, y + h / 2
            cx = g["x"] + (cx - g["chx"]) * g["sx"]
            cy = g["y"] + (cy - g["chy"]) * g["sy"]
            w, h = w * g["sx"], h * g["sy"]
            gcx, gcy = g["x"] + g["w"] / 2, g["y"] + g["h"] / 2
            if g["flipH"]:
                cx = 2 * gcx - cx
                fh, rot = not fh, -rot
            if g["flipV"]:
                cy = 2 * gcy - cy
                fv, rot = not fv, -rot
            if g["rot"]:
                a = math.radians(g["rot"])
                dx, dy = cx - gcx, cy - gcy
                cx, cy = gcx + dx * math.cos(a) - dy * math.sin(a), gcy + dx * math.sin(a) + dy * math.cos(a)
                rot += g["rot"]
            x, y = cx - w / 2, cy - h / 2
        d = self.deck
        rot = ((rot + 180) % 360) - 180
        return {"x": r2(d.x(x)), "y": r2(d.y(y)), "w": r2(max(1, d.px(w))), "h": r2(max(1, d.px(h))),
                **({"rot": r2(rot)} if abs(rot) > 0.01 else {}), **({"flipH": True} if fh else {}), **({"flipV": True} if fv else {})}

    def group_tf(self, grp):
        xfrm = grp.find(f"{P}grpSpPr/{A}xfrm")
        if xfrm is None:
            xfrm = grp.find(f"{DSP}grpSpPr/{A}xfrm")
        if xfrm is None or xfrm.find(A + "off") is None:
            return None
        g = lambda tag, a, b: (int(xfrm.find(A + tag).get(a, 0)), int(xfrm.find(A + tag).get(b, 0))) if xfrm.find(A + tag) is not None else (0, 0)
        (x, y), (w, h) = g("off", "x", "y"), g("ext", "cx", "cy")
        (chx, chy), (chw, chh) = g("chOff", "x", "y"), g("chExt", "cx", "cy")
        return {"x": x, "y": y, "w": w, "h": h, "chx": chx, "chy": chy, "sx": w / chw if chw else 1, "sy": h / chh if chh else 1,
                "rot": int(xfrm.get("rot", 0)) / 60000, "flipH": xfrm.get("flipH") in ("1", "true"), "flipV": xfrm.get("flipV") in ("1", "true")}

    # ---- fills and lines

    def style_ref(self, el, tag):
        ref = el.find(f"{P}style/{A}{tag}") if el.find(P + "style") is not None else el.find(f"{DSP}style/{A}{tag}")
        if ref is None:
            return None, None
        idx = ref.get("idx", "0")
        return (int(idx) if idx.isdigit() else 0), self.color(ref)

    def fill_of(self, sppr, el=None, group_fill=None, ph_chain=()):
        """'none', ('solid', ((r,g,b), alpha)), ('blip', blip element), or None (nothing says)."""
        for node in [sppr, *[n.find(P + "spPr") for n in ph_chain]]:
            if node is None:
                continue
            for child in node:
                name = local(child)
                if name == "noFill":
                    return "none"
                if name == "solidFill":
                    c = self.color(child)
                    return ("solid", c) if c else None
                if name == "gradFill":
                    return self.gradient_of(child)
                if name == "pattFill":
                    self.deck.stats["unsupported"] += 1
                    c = self.color(child.find(A + "fgClr")) or self.color(child.find(A + "bgClr"))
                    return ("solid", c) if c else None
                if name == "blipFill":
                    return ("blip", child)
                if name == "grpFill":
                    return group_fill
        if el is not None:
            idx, color = self.style_ref(el, "fillRef")
            if idx:
                styles = self.theme.bg_fills if idx >= 1000 else self.theme.fills
                i = idx - 1001 if idx >= 1000 else idx - 1
                if 0 <= i < len(styles):
                    st = styles[i]
                    if local(st) == "noFill":
                        return "none"
                    if local(st) == "solidFill":
                        c = color_of(st, self.theme, color, self.clr_map)
                        return ("solid", c) if c else None
                    if local(st) == "gradFill":
                        return self.gradient_of(st, color)
                if color:
                    return ("solid", color)
        return None

    def gradient_of(self, node, ref=None):
        stops = []
        first_color = None
        for gs in node.iter(A + "gs"):
            c = color_of(gs, self.theme, ref, self.clr_map)
            if c:
                first_color = first_color or c
                stops.append({"at": r2(clamp(int(gs.get("pos", 0)) / 100000, 0, 1)),
                              "color": hexc(c), "opacity": r2(clamp(c[1], 0, 1))})
        if not stops:
            return None
        lin = node.find(A + "lin")
        if lin is None:
            self.deck.stats["unsupported"] += 1
            return ("solid", first_color)
        return ("gradient", {"angle": r2(int(lin.get("ang", 0)) / 60000), "stops": stops})

    def shadow_of(self, sppr, el=None, ph_chain=()):
        effect = next((s.find(f"{A}effectLst/{A}outerShdw") for s in [sppr, *[n.find(P + "spPr") for n in ph_chain]]
                       if s is not None and s.find(f"{A}effectLst/{A}outerShdw") is not None), None)
        ref_color = None
        if effect is None and el is not None:
            idx, ref_color = self.style_ref(el, "effectRef")
            if idx and 0 < idx <= len(self.theme.effects):
                effect = self.theme.effects[idx - 1].find(f"{A}effectLst/{A}outerShdw")
        if effect is None:
            return None
        color = self.color(effect, ref_color)
        if not color:
            return None
        distance = self.deck.px(int(effect.get("dist", 0)))
        angle = math.radians(int(effect.get("dir", 0)) / 60000)
        return {"dx": r2(distance * math.cos(angle)), "dy": r2(distance * math.sin(angle)),
                "blur": r2(self.deck.px(int(effect.get("blurRad", 0)))),
                "color": hexc(color), "opacity": r2(clamp(color[1], 0, 1))}

    def line_of(self, sppr, el=None, ph_chain=()):
        """{color, alpha, w, dash, head, tail, headSize, tailSize} in pixels, or None for no line."""
        ln = None
        for node in [sppr, *[n.find(P + "spPr") for n in ph_chain]]:
            if node is not None and node.find(A + "ln") is not None:
                ln = node.find(A + "ln")
                break
        ref_idx, ref_color = self.style_ref(el, "lnRef") if el is not None else (None, None)
        theme_ln = self.theme.lines[ref_idx - 1] if ref_idx and 0 < ref_idx <= len(self.theme.lines) else None
        width = None
        color = None
        explicit_none = False
        dash = None
        for node in (ln, theme_ln):
            if node is None:
                continue
            if width is None and node.get("w"):
                width = int(node.get("w"))
            if color is None and not explicit_none:
                if node.find(A + "noFill") is not None:
                    explicit_none = True
                elif node.find(A + "solidFill") is not None:
                    color = color_of(node.find(A + "solidFill"), self.theme, ref_color, self.clr_map)
                elif node.find(A + "gradFill") is not None:
                    self.deck.stats["unsupported"] += 1
                    stops = [color_of(gs, self.theme, ref_color, self.clr_map) for gs in node.iter(A + "gs")]
                    color = next((s for s in stops if s), None)
            if dash is None and node.find(A + "prstDash") is not None:
                dash = node.find(A + "prstDash").get("val")
        if explicit_none or (color is None and ref_color is None) or (ln is None and theme_ln is None and not ref_idx):
            return None
        color = color or ref_color
        out = {"color": color, "w": max(0.5, self.deck.px(width if width is not None else 9525))}
        if DASHES.get(dash):
            out["dash"] = DASHES[dash]
        for end, key in (("headEnd", "head"), ("tailEnd", "tail")):
            node = ln.find(A + end) if ln is not None else None
            if node is not None and ARROWS.get(node.get("type")):
                out[key] = ARROWS[node.get("type")]
                size = {"sm": 1, "med": 2, "lg": 3}.get(node.get("w") or node.get("len") or "med", 2)
                if size != 2:
                    out[key + "Size"] = size
        return out

    # ---- text

    def lst_chain(self, el, ph_chain, ph_type):
        """List styles from the most specific to the least: shape, layout, master placeholder, master text styles, deck."""
        chain = []
        body = text_body(el)
        if body is not None and body.find(A + "lstStyle") is not None:
            chain.append(body.find(A + "lstStyle"))
        for node in ph_chain:
            lst = node.find(f"{P}txBody/{A}lstStyle")
            if lst is not None:
                chain.append(lst)
        styles = self.master._element.find(f"{P}txStyles")
        if styles is not None:
            if ph_type in ("title", "ctrTitle"):
                chain.append(styles.find(P + "titleStyle"))
            elif ph_type in ("body", "obj", "subTitle"):
                chain.append(styles.find(P + "bodyStyle"))
            elif ph_type:
                chain.append(styles.find(P + "otherStyle"))
        if self.deck.default_text is not None:
            chain.append(self.deck.default_text)
        return [c for c in chain if c is not None]

    def body_pr(self, el, ph_chain):
        out = {}
        for node in (el, *ph_chain):
            body = text_body(node)
            bpr = body.find(A + "bodyPr") if body is not None else None
            if bpr is None:
                continue
            for key in ("lIns", "tIns", "rIns", "bIns", "anchor", "wrap", "vert"):
                if key not in out and bpr.get(key) is not None:
                    out[key] = bpr.get(key)
            if "fontScale" not in out:
                fit = bpr.find(A + "normAutofit")
                if fit is not None:
                    out["fontScale"] = int(fit.get("fontScale", 100000)) / 100000
                    out["lnSpcReduction"] = int(fit.get("lnSpcReduction", 0)) / 100000
        return out

    @staticmethod
    def level_pr(chain, level):
        """pPr elements for a level along the chain (most specific first)."""
        out = []
        for lst in chain:
            for tag in (f"lvl{level + 1}pPr",):
                node = lst.find(A + tag)
                if node is not None:
                    out.append(node)
        return out

    def text_of(self, el, ph_chain, ph_type, *, ref_font_color=None):
        """The words of a shape as rich text with their settings: (html, base settings, box settings) or None."""
        body = text_body(el)
        if body is None:
            return None
        paragraphs = body.findall(A + "p")
        if not any("".join(t.text or "" for t in p.iter(A + "t")).strip() for p in paragraphs):
            return None
        bpr = self.body_pr(el, ph_chain)
        chain = self.lst_chain(el, ph_chain, ph_type)
        own_lst = body.find(A + "lstStyle") is not None and len(body.find(A + "lstStyle"))
        font_scale = bpr.get("fontScale", 1.0)
        spacing_cut = bpr.get("lnSpcReduction", 0.0)
        part = self.slide.part
        counters: dict[int, int] = {}
        html = []
        sizes: list[tuple[float, int]] = []
        first = None
        for p in paragraphs:
            ppr = p.find(A + "pPr")
            level = int(ppr.get("lvl", 0)) if ppr is not None else 0
            levels = ([ppr] if ppr is not None else []) + self.level_pr(chain, level)

            def ppr_get(attr, levels=levels):
                for node in levels:
                    if node.get(attr) is not None:
                        return node.get(attr)
                return None

            def ppr_child(tag, levels=levels):
                for node in levels:
                    child = node.find(A + tag)
                    if child is not None:
                        return child
                return None

            def def_rpr(attr, levels=levels):
                for node in levels:
                    d = node.find(A + "defRPr")
                    if d is not None and d.get(attr) is not None:
                        return d.get(attr)
                return None

            own = len(([ppr] if ppr is not None else []) + (self.level_pr(chain[:1], level) if own_lst else []))

            def def_child(tag, levels=levels, own=own):
                # A shape's style colour (fontRef) comes before what the master and the deck give.
                for i, node in enumerate(levels):
                    if i == own and tag == "solidFill" and ref_font_color is not None:
                        return None
                    d = node.find(A + "defRPr")
                    if d is not None and d.find(A + tag) is not None:
                        return d.find(A + tag)
                return None

            align = {"ctr": "center", "r": "right", "just": "justify", "dist": "distributed", "thaiDist": "distributed"}.get(ppr_get("algn") or "l", "left")
            base_sz = def_rpr("sz")
            base_size = (int(base_sz) / 100 if base_sz else 18.0) * font_scale
            base_bold = def_rpr("b") in ("1", "true")
            base_italic = def_rpr("i") in ("1", "true")
            fill_el = def_child("solidFill")
            base_color = self.color(fill_el) if fill_el is not None else ref_font_color
            latin = def_child("latin")
            ea = def_child("ea")
            base_font = self.theme.font((ea.get("typeface") if ea is not None else None) or (latin.get("typeface") if latin is not None else None))
            # Line spacing: a percentage of the font's line (its ascent + descent), or points.
            ln = ppr_child("lnSpc")
            line = {"pct": 1.0}
            if ln is not None and ln.find(A + "spcPct") is not None:
                line = {"pct": int(ln.find(A + "spcPct").get("val", 100000)) / 100000}
            elif ln is not None and ln.find(A + "spcPts") is not None:
                line = {"pts": int(ln.find(A + "spcPts").get("val", 0)) / 100}
            if "pct" in line and spacing_cut:
                line["pct"] = max(0.5, line["pct"] - spacing_cut)
            spc = []
            for tag in ("spcBef", "spcAft"):
                node = ppr_child(tag)
                if node is not None and node.find(A + "spcPts") is not None:
                    spc.append(int(node.find(A + "spcPts").get("val", 0)) / 100)
                elif node is not None and node.find(A + "spcPct") is not None:
                    spc.append(int(node.find(A + "spcPct").get("val", 0)) / 100000 * base_size)
                else:
                    spc.append(0.0)
            # Bullets and numbers
            bullet = None
            chosen = None
            for node in levels:
                for tag in ("buNone", "buChar", "buAutoNum", "buBlip"):
                    if node.find(A + tag) is not None:
                        chosen = node.find(A + tag)
                        break
                if chosen is not None:
                    break
            if chosen is not None and local(chosen) == "buChar":
                bullet = chosen.get("char", "•")
                bullet = WINGDINGS.get(bullet, bullet)
            elif chosen is not None and local(chosen) == "buAutoNum":
                start = int(chosen.get("startAt", 1))
                counters[level] = counters.get(level, start - 1) + 1
                bullet = number_label(chosen.get("type", "arabicPeriod"), counters[level])
            elif chosen is not None and local(chosen) == "buBlip":
                bullet = "•"
            for deeper in [lv for lv in counters if lv > level]:
                counters.pop(deeper, None)
            if bullet is None:
                counters.pop(level, None)
            # Runs
            runs_html = []
            line_sizes = []
            for node in p:
                name = local(node)
                if name not in ("r", "br", "fld"):
                    continue
                if name == "br":
                    runs_html.append("<br>")
                    continue
                text = "".join(t.text or "" for t in node.iter(A + "t"))
                if name == "fld" and node.get("type") == "slidenum":
                    text = str(self.number)
                if not text:
                    continue
                rpr = node.find(A + "rPr")
                get = lambda attr, rpr=rpr: rpr.get(attr) if rpr is not None and rpr.get(attr) is not None else None
                size = (int(get("sz")) / 100 * font_scale) if get("sz") else base_size
                bold = get("b") in ("1", "true") if get("b") is not None else base_bold
                italic = get("i") in ("1", "true") if get("i") is not None else base_italic
                under = (get("u") or def_rpr("u") or "none") != "none"
                strike = (get("strike") or def_rpr("strike") or "noStrike") != "noStrike"
                baseline = int(get("baseline") or 0)
                fill = rpr.find(A + "solidFill") if rpr is not None else None
                color = self.color(fill) if fill is not None else base_color
                if rpr is not None and rpr.find(A + "noFill") is not None:
                    color = None
                font = None
                if rpr is not None:
                    for tag in ("ea", "latin"):
                        f = rpr.find(A + tag)
                        if f is not None and f.get("typeface"):
                            font = self.theme.font(f.get("typeface"))
                            break
                font = font or base_font
                spc_attr = get("spc")
                cap = get("cap")
                if cap == "all":
                    text = text.upper()
                line_sizes.append(size)
                sizes.append((size, len(text)))
                run = {"size": size, "bold": bold, "italic": italic, "under": under, "strike": strike, "color": color,
                       "font": font, "spc": int(spc_attr) / 100 if spc_attr else 0.0, "baseline": baseline}
                if first is None:
                    first = {**run, "align": align, "line": line, "space": spc, "bullet": bullet}
                link = None
                click = rpr.find(A + "hlinkClick") if rpr is not None else None
                if click is not None and click.get(R + "id"):
                    try:
                        rel = part.rels[click.get(R + "id")]
                        if rel.is_external and re.match(r"^(https?://|mailto:)", rel.target_ref, re.I):
                            link = rel.target_ref
                    except KeyError:
                        link = None
                runs_html.append((run, esc(text).replace("\t", " ").replace("\v", "<br>"), link))
            para_size = max(line_sizes) if line_sizes else base_size
            if first is None and not runs_html:
                # An empty paragraph keeps its line.
                sizes.append((base_size, 0))
            html.append({"align": align, "level": level, "bullet": bullet, "runs": runs_html, "size": para_size, "line": line, "spc": spc, "font": base_font})
        if first is None:
            return None
        # The box's own settings are those of its most used size; runs say what differs.
        base_size = max(set(s for s, _ in sizes), key=lambda s: sum(n for z, n in sizes if z == s)) if sizes else first["size"]
        base = {"size": base_size, "bold": first["bold"], "italic": first["italic"], "color": first["color"], "font": first["font"],
                "under": first["under"], "strike": first["strike"]}
        out = []
        for para in html:
            parts = []
            for item in para["runs"]:
                if item == "<br>":
                    parts.append("<br>")
                    continue
                run, text, link = item
                style = []
                if run["color"] and (not base["color"] or hexc(run["color"]) != hexc(base["color"])):
                    style.append(f"color: {hexc(run['color'])}")
                if abs(run["size"] - base["size"]) > 0.25:
                    style.append(f"font-size: {r2(self.deck.pt(run['size']))}px")
                if run["font"] and run["font"] != base["font"] and re.fullmatch(r"[\w\s.+\-]{1,100}", run["font"]):
                    style.append(f'font-family: "{run["font"]}"')
                if run["bold"] != base["bold"]:
                    style.append(f"font-weight: {700 if run['bold'] else 400}")
                if run["italic"] != base["italic"]:
                    style.append(f"font-style: {'italic' if run['italic'] else 'normal'}")
                deco = [d for d, on, b in (("underline", run["under"], base["under"]), ("line-through", run["strike"], base["strike"])) if on and not b]
                if deco:
                    style.append(f"text-decoration: {' '.join(deco)}")
                html_text = text
                if run["baseline"] > 0:
                    html_text = f"<sup>{html_text}</sup>"
                elif run["baseline"] < 0:
                    html_text = f"<sub>{html_text}</sub>"
                if style:
                    html_text = f'<span style="{"; ".join(style)}">{html_text}</span>'
                if link:
                    html_text = f'<a href="{esc(link)}">{html_text}</a>'
                parts.append(html_text)
            attrs = []
            if para["align"] == "distributed":
                attrs.append('style="text-align: justify; text-align-last: justify"')
            elif para["align"] != "left":
                attrs.append(f'style="text-align: {para["align"]}"')
            if para["level"]:
                attrs.append(f'data-indent="{min(4, para["level"])}"')
            if para["bullet"]:
                attrs.append(f'data-bullet="{esc(para["bullet"][:4])}"')
            inner = "".join(parts) or "<br>"
            out.append(f"<p{(' ' + ' '.join(attrs)) if attrs else ''}>{inner}</p>")
        fs = self.deck.pt(base["size"])
        line = first["line"]
        factor = next((f for pattern, f in LINE_FACTORS if base["font"] and pattern.search(base["font"])), 1.2)
        lh = (line["pct"] * factor) if "pct" in line else (self.deck.pt(line["pts"]) / fs if fs else 1.2)
        before, after = first["space"]
        psp = (before + after) / base["size"] if base["size"] else 0
        ls = first["spc"] / base["size"] if base["size"] else 0
        settings = {
            "fs": r2(fs), "lh": r2(clamp(lh, 0.8, 4)), "align": first["align"],
            **({"color": hexc(base["color"])} if base["color"] else {}),
            **({"bold": True} if base["bold"] else {}), **({"italic": True} if base["italic"] else {}),
            **({"underline": True} if base["under"] else {}), **({"strike": True} if base["strike"] else {}),
            **({"psp": r2(clamp(psp, 0, 4))} if psp > 0.01 else {}),
            **({"ls": r2(clamp(ls, -0.2, 1))} if abs(ls) > 0.005 else {}),
        }
        font = base["font"] or ""
        if font and re.fullmatch(r"[\w\s.+\-]{1,100}", font):
            settings["fontFace"] = font
        if re.search(r"BIZ UD", font):
            settings["font"] = "ud"
        elif re.search(r"教科書|Kyokasho", font):
            settings["font"] = "kyokasho"
        ins = lambda key, default: self.deck.px(int(bpr.get(key, default)))
        box = {
            "pad": [r2(ins("tIns", 45720)), r2(ins("rIns", 91440)), r2(ins("bIns", 45720)), r2(ins("lIns", 91440))],
            "valign": {"ctr": "middle", "b": "bottom"}.get(bpr.get("anchor", "t"), "top"),
            **({"wrap": False} if bpr.get("wrap") == "none" else {}),
            **({"vertical": True} if bpr.get("vert") in ("vert", "eaVert", "wordArtVert", "wordArtVertRtl", "vert270", "mongolianVert") else {}),
        }
        return "".join(out), settings, box

    # ---- shapes

    def add(self, o: dict[str, Any], spid: str | None, group: str | None):
        # Keep off-page objects too. They may be moved into view by an animation or an HTML action.
        if len(self.objects) >= 5000:
            raise ValueError(f"スライド{self.number}の部品が5000個を超えています。")
        if group:
            o["group"] = group
        if spid and not self.importing_decoration and spid not in self.by_spid:
            self.by_spid[spid] = o["id"]
        if spid and not self.importing_decoration and spid in self.actions:
            o["action"] = self.actions.pop(spid)
        self.objects.append(o)

    def shape(self, el, tf, group, *, part=None, offset=None):
        """A p:sp (or dsp:sp): a shape, a text box, or a line."""
        info = self.ph_info(el)
        ph_chain = self.inherited_ph(info) if info else []
        xf = self.xfrm_of(el, ph_chain)
        if xf is None:
            return
        if offset:
            xf = {**xf, "x": xf["x"] + offset[0], "y": xf["y"] + offset[1]}
        nv = el.find(f"{P}nvSpPr/{P}cNvPr")
        if nv is None:
            nv = el.find(f"{DSP}nvSpPr/{DSP}cNvPr")
        spid = nv.get("id") if nv is not None else None
        name = nv.get("name") if nv is not None else None
        hidden = nv is not None and nv.get("hidden") in ("1", "true")
        sppr = el.find(P + "spPr") if el.find(P + "spPr") is not None else el.find(DSP + "spPr")
        if sppr is not None and (sppr.find(A + "scene3d") is not None or sppr.find(A + "sp3d") is not None):
            self.deck.stats["unsupported"] += 1
        b = self.box(xf, tf)
        geom = sppr.find(A + "prstGeom") if sppr is not None else None
        cust = sppr.find(A + "custGeom") if sppr is not None else None
        prst = geom.get("prst") if geom is not None else ("custom" if cust is not None else "rect")
        if prst in ("line", "straightConnector1", "bentConnector2", "bentConnector3", "bentConnector4", "bentConnector5",
                    "curvedConnector2", "curvedConnector3", "curvedConnector4", "curvedConnector5"):
            self.connector(el, xf, tf, group, prst)
            return
        fill = self.fill_of(sppr, el, ph_chain=ph_chain)
        line = self.line_of(sppr, el, ph_chain=ph_chain)
        font_idx, font_color = self.style_ref(el, "fontRef")
        text = self.text_of(el, ph_chain, info["type"] if info else None, ref_font_color=font_color)
        if info and info["type"] in ("title", "ctrTitle") and text and not self.title:
            self.title = re.sub(r"<[^>]+>", " ", text[0]).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
            self.title = re.sub(r"\s+", " ", self.title).strip()
        if fill and fill[0] == "blip":
            # A picture fill: the picture, cut to the shape's outline.
            pic = self.picture_from_blip(fill[1], b, name, part)
            if pic:
                key = SHAPE_MAP.get(prst)
                if key and key != "rect":
                    pic["mask"] = key
                if line:
                    pic.update(self.stroke(line))
                if hidden:
                    pic["hidden"] = True
                self.add(pic, spid, group)
            fill = "none" if text else None
            if not text:
                return
        o = {"id": self.new_id("x", spid), "kind": "shape", **b}
        if name:
            o["name"] = name[:60]
        if cust is not None:
            subs = custom_paths(cust)
            closed_subs = [s for s in subs if len(s["pts"]) > 1]
            if not closed_subs:
                self.deck.stats["unsupported"] += 1
                o.update({"shape": "rect", **self.paint(fill, line)})
                if text:
                    o.update(self.text_settings(text, o))
                if hidden:
                    o["hidden"] = True
                self.add(o, spid, group)
                return
            if len(closed_subs) > 1:
                # Each part of the outline is a drawn shape; they move together.
                group = group or f"g{self.index}c{spid}"
                for k, sub in enumerate(closed_subs):
                    part_o = {"id": self.new_id("x", f"{spid}p{k}"), "kind": "shape", **b, "shape": "custom", "path": {"pts": sub["pts"], **({"closed": True} if sub["closed"] else {})}}
                    part_o.update(self.paint(fill if sub["fill"] else "none", line if sub["stroke"] else None))
                    if k == len(closed_subs) - 1 and text:
                        part_o.update(self.text_settings(text, o))
                    if hidden:
                        part_o["hidden"] = True
                    self.add(part_o, spid, group)
                return
            sub = closed_subs[0]
            o["shape"] = "custom"
            o["path"] = {"pts": sub["pts"], **({"closed": True} if sub["closed"] else {})}
            if not sub["fill"]:
                fill = "none"
            if not sub["stroke"]:
                line = None
        else:
            key = SHAPE_MAP.get(prst)
            if key is None:
                key = "rightArrow" if "Arrow" in prst else "wedgeRectCallout" if "Callout" in prst else "star5" if prst.startswith("star") else "flowProcess" if prst.startswith("flowChart") else "rect"
                self.deck.stats["unsupported"] += 1
            o["shape"] = key
            adj = adjustments(geom, key)
            if adj:
                o["adj"] = adj
        o.update(self.paint(fill, line))
        shadow = self.shadow_of(sppr, el, ph_chain)
        if shadow:
            o["shadow"] = shadow
        if text:
            o.update(self.text_settings(text, o))
            if fill in (None, "none") and not line and o["shape"] == "rect":
                o["kind"] = "text"
                o["fill"] = "none"
        elif fill in (None, "none") and not line:
            return  # nothing to see
        if hidden:
            o["hidden"] = True
        self.add(o, spid, group)

    def paint(self, fill, line):
        out = {}
        if fill and fill not in ("none",) and fill[0] == "solid":
            (rgb, alpha) = fill[1]
            out["fill"] = hexc(fill[1])
            if alpha < 0.995:
                out["fillOpacity"] = r2(clamp(alpha, 0, 1))
        elif fill and fill[0] == "gradient":
            out["fill"] = "none"
            out["gradient"] = fill[1]
        else:
            out["fill"] = "none"
        out.update(self.stroke(line) if line else {"stroke": "none"})
        return out

    def stroke(self, line):
        out = {"stroke": hexc(line["color"]), "strokeW": r2(clamp(line["w"], 0.5, 200))}
        if line.get("dash"):
            out["dash"] = line["dash"]
        return out

    def text_settings(self, text, o):
        html, settings, box = text
        out = {"text": html, **settings, **box, "autofit": "none"}
        if "color" not in out:
            out["color"] = "#000000"
        return out

    def connector(self, el, xf, tf, group, prst):
        nv = el.find(f".//{P}cNvPr")
        spid = nv.get("id") if nv is not None else None
        sppr = el.find(P + "spPr")
        line = self.line_of(sppr, el)
        if not line:
            return
        b = self.box(xf, tf)
        cx, cy = b["x"] + b["w"] / 2, b["y"] + b["h"] / 2
        x1, y1, x2, y2 = b["x"], b["y"], b["x"] + b["w"], b["y"] + b["h"]
        if b.get("flipH"):
            x1, x2 = x2, x1
        if b.get("flipV"):
            y1, y2 = y2, y1
        rot = math.radians(b.get("rot", 0))
        rotate = lambda px, py: (cx + (px - cx) * math.cos(rot) - (py - cy) * math.sin(rot), cy + (px - cx) * math.sin(rot) + (py - cy) * math.cos(rot))
        (x1, y1), (x2, y2) = rotate(x1, y1), rotate(x2, y2)
        o = {"id": self.new_id("l", spid), "kind": "line", "x1": r2(x1), "y1": r2(y1), "x2": r2(x2), "y2": r2(y2),
             "stroke": hexc(line["color"]), "strokeW": r2(clamp(line["w"], 0.5, 200))}
        if line["color"][1] < 0.995:
            o["opacity"] = r2(line["color"][1])
        for key in ("dash", "head", "tail", "headSize", "tailSize"):
            if line.get(key):
                o[key] = line[key]
        if prst.startswith("bent"):
            o["route"] = "elbow"
        elif prst.startswith("curved"):
            o["route"] = "curve"
        nv = el.find(f".//{P}cNvPr")
        if nv is not None and nv.get("hidden") in ("1", "true"):
            o["hidden"] = True
        self.add(o, spid, group)

    def picture_from_blip(self, blip_fill, b, name=None, part=None):
        blip = blip_fill.find(A + "blip")
        if blip is None:
            return None
        rid = blip.get(R + "embed")
        target = self.rel_target(part or self.slide.part, rid) if rid else None
        if target is None:
            return None
        url = image_url(target, self.deck)
        if not url:
            self.deck.stats["unsupported"] += 1
            return None
        o = {"id": self.new_id("i", None), "kind": "image", **b, "src": url, "fit": "fill"}
        if name:
            o["name"] = name[:60]
        rect = blip_fill.find(A + "srcRect")
        if rect is not None:
            crop = {k: max(0.0, int(rect.get(k, 0)) / 100000) for k in ("l", "t", "r", "b")}
            if any(crop.values()) and crop["l"] + crop["r"] < 0.98 and crop["t"] + crop["b"] < 0.98:
                o["crop"] = {k: round(v, 4) for k, v in crop.items()}
        for mod in blip:
            if local(mod) == "alphaModFix":
                o["opacity"] = r2(int(mod.get("amt", 100000)) / 100000)
            if local(mod) == "grayscl":
                o["gray"] = True
        self.deck.stats["pictures"] += 1
        return o

    def picture(self, el, tf, group, part=None, offset=None):
        info = self.ph_info(el)
        xf = self.xfrm_of(el, self.inherited_ph(info) if info else [])
        if xf is None:
            return
        if offset:
            xf = {**xf, "x": xf["x"] + offset[0], "y": xf["y"] + offset[1]}
        nv = el.find(f"{P}nvPicPr/{P}cNvPr")
        spid = nv.get("id") if nv is not None else None
        hidden = nv is not None and nv.get("hidden") in ("1", "true")
        b = self.box(xf, tf)
        o = self.picture_from_blip(el.find(P + "blipFill"), b, nv.get("name") if nv is not None else None, part)
        if not o:
            return
        o["id"] = self.new_id("i", spid)
        sppr = el.find(P + "spPr")
        geom = sppr.find(A + "prstGeom") if sppr is not None else None
        key = SHAPE_MAP.get(geom.get("prst")) if geom is not None else None
        if key and key != "rect":
            o["mask"] = key
            adj = adjustments(geom, key)
            if adj:
                o["adj"] = adj
        line = self.line_of(sppr, el)
        if line:
            o.update(self.stroke(line))
        shadow = self.shadow_of(sppr, el)
        if shadow:
            o["shadow"] = shadow
        if hidden:
            o["hidden"] = True
        self.add(o, spid, group)

    def model3d(self, data, b, spid, part, nv):
        """A 3D model: its GLB file, how PowerPoint turned it, and the picture PowerPoint drew of it (the poster)."""
        m = data.find(AM3D + "model3d")
        rid = m.get(R + "embed") if m is not None else None
        target = self.rel_target(part or self.slide.part, rid) if rid else None
        if target is None or not target.blob or len(target.blob) > 100_000_000:
            self.deck.stats["unsupported"] += 1
            return None
        o = {"id": self.new_id("m", spid), "kind": "model", **b, "src": "data:model/gltf-binary;base64," + base64.b64encode(target.blob).decode()}
        name = nv.get("name") if nv is not None else None
        if name:
            o["name"] = name[:60]
        if nv is not None and nv.get("descr"):
            o["alt"] = nv.get("descr")[:500]
        # The turn (60000ths of a degree): across the model (ax, a tilt), round its upright (ay) and in the slide (az).
        rot = m.find(f"{AM3D}trans/{AM3D}rot")
        if rot is not None:
            deg = lambda k: round((int(rot.get(k, "0")) / 60000 + 180) % 360 - 180, 1)
            view = {"pitch": max(-89, min(89, deg("ax"))), "yaw": deg("ay"), "roll": deg("az")}
            view = {k: v for k, v in view.items() if v}
            if view:
                o["view"] = view
        blip = m.find(f"{AM3D}raster/{AM3D}blip")
        poster_part = self.rel_target(part or self.slide.part, blip.get(R + "embed")) if blip is not None and blip.get(R + "embed") else None
        poster = image_url(poster_part, self.deck) if poster_part is not None else None
        if poster and poster.startswith(("data:image/png", "data:image/jpeg", "data:image/webp")):
            o["poster"] = poster
        self.deck.stats["models"] += 1
        return o

    def frame(self, el, tf, group, part=None):
        """A graphic frame: a table, a chart, SmartArt, or an embedded object (its picture)."""
        xf = self.xfrm_of(el)
        if xf is None:
            return
        nv = el.find(f"{P}nvGraphicFramePr/{P}cNvPr")
        spid = nv.get("id") if nv is not None else None
        b = self.box(xf, tf)
        data = el.find(f"{A}graphic/{A}graphicData")
        if data is None:
            return
        uri = data.get("uri", "")
        if uri.endswith("/table"):
            table = self.table(data.find(A + "tbl"), b, spid)
            if table:
                self.add(table, spid, group)
        elif uri.endswith("/chart"):
            chart = self.chart(data, b, spid, part)
            if chart:
                self.add(chart, spid, group)
        elif uri == NS["cx"]:
            chart = self.chartex(data, b, spid, part)
            if chart:
                self.add(chart, spid, group)
        elif uri.endswith("/diagram"):
            self.smartart(data, xf, tf, group, spid, part)
        elif uri == NS["am3d"]:
            model = self.model3d(data, b, spid, part, nv)
            if model:
                self.add(model, spid, group)
        else:
            # An embedded object (Excel, a picture of an equation…): the picture PowerPoint keeps of it.
            pic = next(iter(data.iter(P + "pic")), None)
            if pic is not None:
                blip_fill = pic.find(P + "blipFill")
                o = self.picture_from_blip(blip_fill, b, part=part) if blip_fill is not None else None
                if o:
                    o["id"] = self.new_id("i", spid)
                    self.add(o, spid, group)
                    return
            self.deck.stats["unsupported"] += 1

    def table(self, tbl, b, spid):
        if tbl is None:
            return None
        cols = [int(c.get("w", 0)) for c in tbl.findall(f"{A}tblGrid/{A}gridCol")]
        rows = tbl.findall(A + "tr")
        if not cols or not rows:
            return None
        if len(rows) > 500 or len(cols) > 100:
            raise ValueError(f"スライド{self.number}の表が対応サイズ（500行・100列）を超えています。")
        tblpr = tbl.find(A + "tblPr")
        flags = {k: tblpr is not None and tblpr.get(k) in ("1", "true") for k in ("firstRow", "lastRow", "firstCol", "lastCol", "bandRow", "bandCol")}
        style_id = tblpr.find(A + "tableStyleId").text if tblpr is not None and tblpr.find(A + "tableStyleId") is not None else None
        style = self.deck.table_styles.get(style_id)
        heights = [int(r.get("h", 0)) for r in rows]
        cells = []
        sizes = []
        n_rows, n_cols = len(rows), len(cols)
        for ri, tr in enumerate(rows):
            row = []
            for ci, tc in enumerate(tr.findall(A + "tc")[:n_cols]):
                if tc.get("hMerge") in ("1", "true") or tc.get("vMerge") in ("1", "true"):
                    row.append({"merged": True})
                    continue
                cell: dict[str, Any] = {}
                parts = table_parts(flags, ri, ci, n_rows, n_cols)
                fill, txt_bold, txt_color, borders = None, None, None, {}
                for part_name in parts:
                    node = style.find(A + part_name) if style is not None else None
                    if node is None:
                        continue
                    tcs = node.find(A + "tcStyle")
                    if tcs is not None:
                        f = tcs.find(A + "fill")
                        if f is not None and len(f):
                            if f.find(A + "noFill") is not None:
                                fill = "none"
                            elif f.find(A + "solidFill") is not None:
                                fill = self.color(f.find(A + "solidFill"))
                        bd = tcs.find(A + "tcBdr")
                        if bd is not None:
                            r0, c0, r1, c1 = part_region(part_name, ri, ci, n_rows, n_cols)
                            for side in bd:
                                ln = side.find(A + "ln")
                                if ln is None:
                                    continue
                                name = local(side)
                                if name == "left" and ci == c0 or name == "right" and ci == c1 or name == "top" and ri == r0 or name == "bottom" and ri == r1:
                                    borders[name] = ln
                                elif name == "insideH":
                                    if ri < r1:
                                        borders["bottom"] = ln
                                    if ri > r0:
                                        borders["top"] = ln
                                elif name == "insideV":
                                    if ci < c1:
                                        borders["right"] = ln
                                    if ci > c0:
                                        borders["left"] = ln
                    tts = node.find(A + "tcTxStyle")
                    if tts is not None:
                        if tts.get("b") == "on":
                            txt_bold = True
                        c = self.color(tts)
                        if c is None and tts.find(A + "fontRef") is not None:
                            c = self.color(tts.find(A + "fontRef"))
                        if c:
                            txt_color = c
                tcpr = tc.find(A + "tcPr")
                if tcpr is not None:
                    if tcpr.find(A + "noFill") is not None:
                        fill = "none"
                    elif tcpr.find(A + "solidFill") is not None:
                        fill = self.color(tcpr.find(A + "solidFill"))
                    for tag, side in (("lnL", "left"), ("lnR", "right"), ("lnT", "top"), ("lnB", "bottom")):
                        ln = tcpr.find(A + tag)
                        if ln is not None:
                            borders[side] = ln
                if fill and fill != "none":
                    cell["fill"] = hexc(fill)
                for side, ln in borders.items():
                    key = {"left": "bl", "right": "br", "top": "bt", "bottom": "bb", "insideH": None, "insideV": None}.get(side)
                    if key is None:
                        continue
                    if ln.find(A + "noFill") is not None:
                        cell[key] = "none"
                        continue
                    c = self.color(ln.find(A + "solidFill")) if ln.find(A + "solidFill") is not None else None
                    if c:
                        cell[key] = {"c": hexc(c), "w": r2(max(0.5, self.deck.px(int(ln.get("w", 12700)))))}
                txt = self.text_of(tc, [], None, ref_font_color=txt_color)
                if txt:
                    html, settings, box = txt
                    cell["text"] = html
                    sizes.append(settings["fs"])
                    cell["fs"] = settings["fs"]
                    cell["lh"] = settings["lh"]
                    if settings.get("fontFace"):
                        cell["fontFace"] = settings["fontFace"]
                    if box.get("pad"):
                        cell["pad"] = box["pad"]
                    if settings.get("color") or txt_color:
                        cell["color"] = settings.get("color") or hexc(txt_color)
                    if settings.get("bold") or txt_bold:
                        cell["bold"] = True
                    for key in ("italic", "underline", "strike"):
                        if settings.get(key):
                            cell[key] = True
                    if settings["align"] != "left":
                        cell["align"] = settings["align"]
                anchor = tcpr.get("anchor") if tcpr is not None else None
                cell["valign"] = {"ctr": "middle", "b": "bottom"}.get(anchor or "t", "top")
                # 文字列の方向: vertical text in the cell (縦書き).
                if tcpr is not None and tcpr.get("vert") in ("vert", "eaVert", "wordArtVertRtl"):
                    cell["vertical"] = True
                span, rspan = int(tc.get("gridSpan", 1)), int(tc.get("rowSpan", 1))
                if span > 1:
                    cell["cs"] = span
                if rspan > 1:
                    cell["rs"] = rspan
                row.append(cell)
            while len(row) < n_cols:
                row.append({})
            cells.append(row)
        self.deck.stats["tables"] += 1
        fs = max(set(sizes), key=sizes.count) if sizes else self.deck.pt(18)
        return {"id": self.new_id("t", spid), "kind": "table", **b, "cells": cells, "cols": shares(cols), "rows": shares(heights),
                "style": "plain", "header": False, "banded": False, "fs": r2(fs), "color": "#000000", "lh": 1.2}

    def chart(self, data, b, spid, owner_part=None):
        ref = data.find(f"{{{NS['c']}}}chart")
        rid = ref.get(R + "id") if ref is not None else None
        part = self.rel_target(owner_part or self.slide.part, rid) if rid else None
        if part is None:
            return None
        try:
            root = etree.fromstring(part.blob)
            out = ChartReader(self, root).read()
        except ValueError:
            raise
        except Exception:
            self.deck.stats["unsupported"] += 1
            return None
        if not out:
            return None
        self.deck.stats["charts"] += 1
        return {"id": self.new_id("c", spid), "kind": "chart", **b, "chart": out}

    def chartex(self, data, b, spid, owner_part=None):
        """Office 2016's charts (ツリーマップ・サンバースト・箱ひげ図・ヒストグラム・じょうご・ウォーターフォール) as the studio's own
        editable charts."""
        part = self.chartex_part(data, owner_part)
        if part is None:
            return None
        try:
            out = read_chartex(etree.fromstring(part.blob))
        except Exception:
            out = None
        if not out:
            return None
        self.deck.stats["charts"] += 1
        return {"id": self.new_id("c", spid), "kind": "chart", **b, "chart": out}

    def chartex_part(self, data, owner_part=None):
        ref = data.find(f"{{{NS['cx']}}}chart")
        rid = ref.get(R + "id") if ref is not None else None
        return self.rel_target(owner_part or self.slide.part, rid) if rid else None

    def readable_choice(self, choice, owner_part=None):
        """An mc:Choice this reader understands; otherwise the mc:Fallback (PowerPoint's picture of it) is taken. A
        chart of a kind the studio does not draw (a map…) keeps its picture."""
        for frame in choice.iter(P + "graphicFrame"):
            data = frame.find(f"{A}graphic/{A}graphicData")
            if data is None or data.get("uri") != NS["cx"]:
                continue
            part = self.chartex_part(data, owner_part)
            try:
                if part is None or not read_chartex(etree.fromstring(part.blob)):
                    return False
            except Exception:
                return False
        return True

    def smartart(self, data, xf, tf, group, spid, owner_part=None):
        rel_ids = data.find(f"{{{NS['dgm']}}}relIds")
        if rel_ids is None:
            return
        dm = self.rel_target(owner_part or self.slide.part, rel_ids.get(R + "dm"))
        drawing = None
        if dm is not None:
            root = etree.fromstring(dm.blob)
            ext = root.find(f".//{{{NS['dsp']}}}dataModelExt")
            if ext is not None and ext.get("relId"):
                drawing = self.rel_target(dm, ext.get("relId"))
        if drawing is None:
            self.deck.stats["unsupported"] += 1
            return
        root = etree.fromstring(drawing.blob)
        tree = root.find(f"{DSP}spTree")
        if tree is None:
            return
        group = group or f"g{self.index}d{spid}"
        offset = (xf["x"], xf["y"])
        for child in tree:
            name = local(child)
            if name == "sp":
                self.shape(child, tf, group, part=drawing, offset=offset)
            elif name == "grpSp":
                self.walk(child, tf, group, part=drawing, offset=offset)

    def click_action(self, el):
        """What clicking the shape does in a slide show: go to a slide, next / previous / first / last, end, or open a link."""
        nv = next((c for c in el if local(c).startswith("nv")), None)
        cnv = nv.find(P + "cNvPr") if nv is not None else None
        click = cnv.find(A + "hlinkClick") if cnv is not None else None
        if click is None:
            return None
        action = click.get("action") or ""
        jump = re.match(r"^ppaction://hlinkshowjump\?jump=(\w+)", action)
        if jump:
            kind = {"nextslide": "next", "previousslide": "prev", "firstslide": "first", "lastslide": "last", "endshow": "end"}.get(jump.group(1).lower())
            return {"type": kind} if kind else None
        try:
            rel = self.slide.part.rels[click.get(R + "id")] if click.get(R + "id") else None
        except KeyError:
            rel = None
        if rel is None:
            return None
        if action.startswith("ppaction://hlinksldjump"):
            sid = self.deck.sids.get(str(rel.target_part.partname)) if not rel.is_external else None
            return {"type": "slide", "to": sid} if sid else None
        if not action and rel.is_external and re.match(r"^(https?://|mailto:)", rel.target_ref, re.I):
            return {"type": "url", "href": rel.target_ref[:2000]}
        return None

    def walk(self, tree, tf=(), group=None, part=None, offset=None, decorations=False):
        for child in tree:
            name = local(child)
            if decorations and name in ("sp", "pic", "graphicFrame") and child.find(f".//{P}nvPr/{P}ph") is not None:
                continue
            if part is None and name in ("sp", "pic", "cxnSp", "graphicFrame"):
                action = self.click_action(child)
                nv = next((c for c in child if local(c).startswith("nv")), None)
                cnv = nv.find(P + "cNvPr") if nv is not None else None
                if action and cnv is not None and cnv.get("id"):
                    self.actions[cnv.get("id")] = action
            if name == "AlternateContent":
                choice = child.find(f"{{{NS['mc']}}}Choice")
                fallback = child.find(f"{{{NS['mc']}}}Fallback")
                inner = choice if choice is not None and len(choice) and self.readable_choice(choice, part) else fallback
                if inner is not None:
                    self.walk(inner, tf, group, part, offset, decorations)
                continue
            if name == "sp":
                self.shape(child, tf, group, part=part, offset=offset)
            elif name == "pic":
                self.picture(child, tf, group, part, offset)
            elif name == "cxnSp":
                xf = self.xfrm_of(child)
                if xf is not None:
                    if offset:
                        xf = {**xf, "x": xf["x"] + offset[0], "y": xf["y"] + offset[1]}
                    geom = child.find(f"{P}spPr/{A}prstGeom")
                    self.connector(child, xf, tf, group, geom.get("prst") if geom is not None else "line")
            elif name == "graphicFrame":
                start = len(self.objects)
                self.frame(child, tf, group, part)
                nv = child.find(f"{P}nvGraphicFramePr/{P}cNvPr")
                if nv is not None and nv.get("hidden") in ("1", "true"):
                    for object_in_frame in self.objects[start:]:
                        object_in_frame["hidden"] = True
            elif name == "grpSp":
                g = self.group_tf(child)
                if g is None:
                    continue
                nv = child.find(f"{P}nvGrpSpPr/{P}cNvPr")
                if nv is None:
                    nv = child.find(f"{DSP}nvGrpSpPr/{DSP}cNvPr")
                group_hidden = nv is not None and nv.get("hidden") in ("1", "true")
                gid = group or f"g{self.index}x{nv.get('id') if nv is not None else len(self.objects)}"
                if offset and not tf:
                    g = {**g, "x": g["x"] + offset[0], "y": g["y"] + offset[1]}
                start = len(self.objects)
                self.walk(child, (g, *tf), gid, part, decorations=decorations)
                if group_hidden:
                    for object_in_group in self.objects[start:]:
                        object_in_group["hidden"] = True
                if not self.importing_decoration and nv is not None and nv.get("id") and len(self.objects) > start:
                    self.by_spid.setdefault(nv.get("id"), f"grp:{gid}")

    # ---- backgrounds, layouts and masters

    def background(self):
        for owner in (self.slide, self.layout, self.master):
            bg = owner._element.find(f"{P}cSld/{P}bg")
            if bg is None:
                continue
            bgpr = bg.find(P + "bgPr")
            fill = None
            if bgpr is not None:
                fill = self.fill_of(bgpr)
            else:
                ref = bg.find(P + "bgRef")
                if ref is not None:
                    idx, color = int(ref.get("idx", 0)), self.color(ref)
                    lst = self.theme.bg_fills if idx >= 1000 else self.theme.fills
                    i = idx - 1001 if idx >= 1000 else idx - 1
                    st = lst[i] if 0 <= i < len(lst) else None
                    if st is not None and local(st) == "solidFill":
                        c = color_of(st, self.theme, color, self.clr_map)
                        fill = ("solid", c) if c else None
                    elif st is not None and local(st) == "gradFill":
                        fill = self.gradient_of(st, color)
                    elif st is not None and local(st) == "blipFill":
                        fill = ("blip", st)
                    elif color:
                        fill = ("solid", color)
            if fill is None:
                return
            full = {"x": r2(self.deck.ox), "y": r2(self.deck.oy), "w": r2(self.deck.cx * self.deck.k), "h": r2(self.deck.cy * self.deck.k)}
            if fill[0] == "blip":
                pic = self.picture_from_blip(fill[1], full, "背景", owner.part)
                if pic:
                    pic["name"] = "背景"
                    pic["locked"] = True
                    self.objects.append(pic)
                return
            if fill[0] == "solid":
                (rgb, alpha) = fill[1]
                if rgb == (255, 255, 255) and alpha > 0.99:
                    return
                self.objects.append({"id": self.new_id("bg", None), "kind": "shape", "name": "背景", "shape": "rect", **full, "fill": hexc(fill[1]), "stroke": "none", "locked": True})
            elif fill[0] == "gradient":
                self.objects.append({"id": self.new_id("bg", None), "kind": "shape", "name": "背景", "shape": "rect", **full,
                                     "fill": "none", "gradient": fill[1], "stroke": "none", "locked": True})
            return

    def decorations(self):
        """Another template's master and layout pictures and shapes (not their placeholders), under the slide's."""
        show = lambda el: el.find(f"{P}cSld") is not None and el.get("showMasterSp") not in ("0", "false")
        layers = []
        if show(self.slide._element) and show(self.layout._element):
            layers.append(self.master)
        if show(self.slide._element):
            layers.append(self.layout)
        for owner in layers:
            tree = owner._element.find(f"{P}cSld/{P}spTree")
            self.importing_decoration = True
            try:
                self.walk(tree, part=owner.part, decorations=True)
            finally:
                self.importing_decoration = False

    # ---- the slide

    def read(self) -> dict[str, Any]:
        self.background()
        self.decorations()
        tree = self.slide._element.find(f"{P}cSld/{P}spTree")
        self.walk(tree)
        slide: dict[str, Any] = {"type": "blank", "title": self.title[:90] or f"スライド {self.number}", "hideTitle": True,
                                 "sourceViewport": {"x": r2(self.deck.ox), "y": r2(self.deck.oy),
                                                    "w": r2(self.deck.cx * self.deck.k), "h": r2(self.deck.cy * self.deck.k)}}
        sid = self.deck.sids.get(str(self.slide.part.partname))
        if sid:
            slide["sid"] = sid  # what links to this slide point at
        slide["master"] = "source"
        notes = notes_text(self.slide)
        if notes:
            slide["notes"] = notes
        if self.objects:
            slide["elements"] = self.objects
        timeline = timeline_of(self.slide._element, self.by_spid, self.deck)
        if timeline:
            slide["timeline"] = timeline
            self.deck.stats["animations"] += len(timeline)
        slide.update(transition_of(self.slide._element))
        comments = comments_of(self.slide, self.deck.comment_authors, self.by_spid, self.index)
        if comments:
            slide["comments"] = comments
            self.deck.stats["comments"] += len(comments)
        self.deck.stats["objects"] += len(self.objects)
        return slide


# ---------------------------------------------------------------- review comments

def comment_authors(prs) -> dict[str, str]:
    """Who wrote the comments: classic (commentAuthors.xml, numeric ids) and modern (authors.xml, GUIDs)."""
    authors: dict[str, str] = {}
    for part in prs.part.package.iter_parts():
        if str(part.partname) not in {"/ppt/commentAuthors.xml", "/ppt/authors.xml"}:
            continue
        try:
            root = etree.fromstring(part.blob)
        except etree.XMLSyntaxError:
            continue
        for node in root:
            if node.get("id") is not None:
                authors[str(node.get("id"))] = (node.get("name") or node.get("initials") or "")[:40]
    return authors


def _comment_text(node) -> str:
    body = node.find("{*}txBody")
    if body is not None:  # modern: <p188:txBody>…<a:t>, a paragraph per line
        lines = ["".join(t.text or "" for t in p.iter(A + "t")) for p in body.iter(A + "p")]
        return "\n".join(lines).strip()
    text = node.find("{*}text")  # classic: <p:text>
    return (text.text or "").strip() if text is not None else ""


def comments_of(slide, authors: dict[str, str], by_spid: dict[str, str], index: int) -> list[dict[str, Any]]:
    """The slide's review comments as threads: who and when, replies, resolved, and the shape a modern comment is on."""
    threads: list[dict[str, Any]] = []
    classic: dict[tuple[str, str], dict[str, Any]] = {}
    n = 0
    for rel in slide.part.rels.values():
        if rel.is_external or not rel.reltype.endswith("/relationships/comments"):
            continue
        try:
            root = etree.fromstring(rel.target_part.blob)
        except (etree.XMLSyntaxError, AttributeError):
            continue
        for cm in root:
            if not isinstance(cm.tag, str) or local(cm) != "cm":
                continue
            text = _comment_text(cm)
            if not text:
                continue
            n += 1
            item: dict[str, Any] = {"id": f"pc{index}x{n}", "text": text[:2000]}
            by = authors.get(str(cm.get("authorId")), "")
            if by:
                item["by"] = by
            at = cm.get("created") or cm.get("dt")
            if at:
                item["at"] = at[:40]
            if (cm.get("status") or "").lower() == "resolved":
                item["done"] = True
            mark = next((m for m in cm.iter() if isinstance(m.tag, str) and local(m) == "spMk"), None)
            if mark is not None and by_spid.get(str(mark.get("id"))):
                item["anchor"] = by_spid[str(mark.get("id"))][:40]
            replies = []
            for k, reply in enumerate(r for r in cm.iter() if isinstance(r.tag, str) and local(r) == "reply"):
                rtext = _comment_text(reply)
                if not rtext:
                    continue
                r_item: dict[str, Any] = {"id": f"{item['id']}r{k}", "text": rtext[:2000]}
                if authors.get(str(reply.get("authorId"))):
                    r_item["by"] = authors[str(reply.get("authorId"))]
                if reply.get("created"):
                    r_item["at"] = reply.get("created")[:40]
                replies.append(r_item)
            # A classic reply is a comment of its own that names its parent (p15:threadingInfo / p15:parentCm).
            parent = next((m for m in cm.iter() if isinstance(m.tag, str) and local(m) == "parentCm"), None)
            if parent is not None and (str(parent.get("authorId")), str(parent.get("idx"))) in classic:
                host = classic[(str(parent.get("authorId")), str(parent.get("idx")))]
                host.setdefault("replies", []).append({k: v for k, v in item.items() if k in ("id", "text", "by", "at")})
                continue
            if replies:
                item["replies"] = replies[:100]
            if cm.get("idx") is not None:
                classic[(str(cm.get("authorId")), str(cm.get("idx")))] = item
            threads.append(item)
    return threads[:200]


# ---------------------------------------------------------------- pieces

DEFAULT_TABLE_STYLE = "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"
_LN = '<a:ln w="{w}"><a:solidFill><a:schemeClr val="lt1"/></a:solidFill></a:ln>'
_ACCENT_TX = '<a:tcTxStyle b="on"><a:schemeClr val="lt1"/></a:tcTxStyle>'
MEDIUM_STYLE_2 = (
    f'<a:tblStyle xmlns:a="{NS["a"]}" styleId="{DEFAULT_TABLE_STYLE}">'
    '<a:wholeTbl><a:tcTxStyle><a:schemeClr val="dk1"/></a:tcTxStyle><a:tcStyle><a:tcBdr>'
    + "".join(f"<a:{side}>{_LN.format(w=12700)}</a:{side}>" for side in ("left", "right", "top", "bottom", "insideH", "insideV"))
    + '</a:tcBdr><a:fill><a:solidFill><a:schemeClr val="accent1"><a:tint val="20000"/></a:schemeClr></a:solidFill></a:fill></a:tcStyle></a:wholeTbl>'
    '<a:band1H><a:tcStyle><a:fill><a:solidFill><a:schemeClr val="accent1"><a:tint val="40000"/></a:schemeClr></a:solidFill></a:fill></a:tcStyle></a:band1H>'
    '<a:band1V><a:tcStyle><a:fill><a:solidFill><a:schemeClr val="accent1"><a:tint val="40000"/></a:schemeClr></a:solidFill></a:fill></a:tcStyle></a:band1V>'
    + "".join(f'<a:{part}>{_ACCENT_TX}<a:tcStyle><a:fill><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:fill></a:tcStyle></a:{part}>' for part in ("lastCol", "firstCol"))
    + f'<a:lastRow>{_ACCENT_TX}<a:tcStyle><a:tcBdr><a:top>{_LN.format(w=38100)}</a:top></a:tcBdr><a:fill><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:fill></a:tcStyle></a:lastRow>'
    + f'<a:firstRow>{_ACCENT_TX}<a:tcStyle><a:tcBdr><a:bottom>{_LN.format(w=38100)}</a:bottom></a:tcBdr><a:fill><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:fill></a:tcStyle></a:firstRow>'
    '</a:tblStyle>')


def part_region(part: str, r: int, c: int, n_rows: int, n_cols: int):
    """The cells a table-style part covers (r0, c0, r1, c1): its outer borders are that region's edges."""
    if part in ("firstRow",):
        return 0, 0, 0, n_cols - 1
    if part == "lastRow":
        return n_rows - 1, 0, n_rows - 1, n_cols - 1
    if part == "firstCol":
        return 0, 0, n_rows - 1, 0
    if part == "lastCol":
        return 0, n_cols - 1, n_rows - 1, n_cols - 1
    if part in ("band1H", "band2H"):
        return r, 0, r, n_cols - 1
    if part in ("band1V", "band2V"):
        return 0, c, n_rows - 1, c
    return 0, 0, n_rows - 1, n_cols - 1


def text_body(el):
    """A shape's text: p:txBody on a slide, a:txBody in a table cell, dsp:txBody in SmartArt."""
    for tag in (P + "txBody", A + "txBody", DSP + "txBody"):
        body = el.find(tag)
        if body is not None:
            return body
    return None


def shares(values):
    total = sum(values) or 1
    return [round(v / total, 5) for v in values]


def number_label(kind: str, n: int) -> str:
    def roman(v):
        out = ""
        for value, sym in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
            while v >= value:
                out, v = out + sym, v - value
        return out
    alpha = lambda v: "".join(chr(ord("a") + (v - 1) % 26) for _ in range(1 + (v - 1) // 26))
    core = (roman(n) if "roman" in kind.lower() else alpha(n) if kind.lower().startswith("alpha") else
            "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"[n - 1] if "circleNum" in kind and n <= 20 else str(n))
    if "Lc" in kind:
        core = core.lower()
    elif "Uc" in kind:
        core = core.upper()
    if "circleNum" in kind:
        return core
    if kind.endswith("ParenBoth"):
        return f"({core})"
    if kind.endswith("ParenR"):
        return f"{core})"
    if kind.endswith("Plain"):
        return core
    if kind.endswith("Period"):
        return f"{core}."
    return f"{core}."


def adjustments(geom, key):
    if geom is None:
        return None
    gds = [gd for gd in geom.iter(A + "gd")]
    if not gds:
        return None
    values = []
    for gd in gds:
        m = re.match(r"val\s+(-?\d+)", gd.get("fmla", ""))
        if not m:
            return None
        values.append(int(m.group(1)))
    if key in ADJ_ANGLES:
        return [v / 60000 if is_angle else v / 100000 for v, is_angle in zip(values, ADJ_ANGLES[key])]
    scale = ADJ_SCALE.get(key)
    if scale is None:
        return None
    count = ADJ_COUNT.get(key, len(values))
    return [round(v / 100000 * scale, 4) for v in values[:count]]


def custom_paths(cust) -> list[dict[str, Any]]:
    """A custom outline as point lists (fractions of the shape's box), curves sampled."""
    out = []
    for path in cust.iter(A + "path"):
        pw, ph = int(path.get("w", 0)) or 1, int(path.get("h", 0)) or 1
        fill = path.get("fill", "norm") != "none"
        stroke = path.get("stroke", "1") not in ("0", "false")
        pts: list[list[float]] = []
        cur = [0.0, 0.0]

        def flush(closed=False):
            nonlocal pts
            if len(pts) > 1:
                dedup = [pts[0]] + [p for i, p in enumerate(pts[1:], 1) if abs(p[0] - pts[i - 1][0]) > 1e-6 or abs(p[1] - pts[i - 1][1]) > 1e-6]
                if len(dedup) > 5000:
                    raise ValueError("図形の輪郭が5000点を超えています。")
                out.append({"pts": [[round(x / pw, 4), round(y / ph, 4)] for x, y in dedup][:5000], "closed": closed, "fill": fill and closed, "stroke": stroke})
            pts = []

        def pt_of(el):
            return [float(el.get("x", 0)), float(el.get("y", 0))]

        for cmd in path:
            name = local(cmd)
            if name == "moveTo":
                flush()
                cur = pt_of(cmd.find(A + "pt"))
                pts = [cur]
            elif name == "lnTo":
                cur = pt_of(cmd.find(A + "pt"))
                pts.append(cur)
            elif name in ("cubicBezTo", "quadBezTo"):
                ctrl = [pt_of(p) for p in cmd.findall(A + "pt")]
                p0 = cur
                for i in range(1, 13):
                    t = i / 12
                    if name == "cubicBezTo" and len(ctrl) == 3:
                        c1, c2, p3 = ctrl
                        x = (1 - t) ** 3 * p0[0] + 3 * (1 - t) ** 2 * t * c1[0] + 3 * (1 - t) * t ** 2 * c2[0] + t ** 3 * p3[0]
                        y = (1 - t) ** 3 * p0[1] + 3 * (1 - t) ** 2 * t * c1[1] + 3 * (1 - t) * t ** 2 * c2[1] + t ** 3 * p3[1]
                    elif len(ctrl) == 2:
                        c1, p3 = ctrl
                        x = (1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * c1[0] + t ** 2 * p3[0]
                        y = (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * c1[1] + t ** 2 * p3[1]
                    else:
                        break
                    pts.append([x, y])
                cur = pts[-1] if pts else cur
            elif name == "arcTo":
                wr, hr = float(cmd.get("wR", 0)), float(cmd.get("hR", 0))
                st, sw = float(cmd.get("stAng", 0)) / 60000, float(cmd.get("swAng", 0)) / 60000
                a0 = math.radians(st)
                cx, cy = cur[0] - wr * math.cos(a0), cur[1] - hr * math.sin(a0)
                steps = max(4, int(abs(sw) / 10))
                for i in range(1, steps + 1):
                    a = math.radians(st + sw * i / steps)
                    pts.append([cx + wr * math.cos(a), cy + hr * math.sin(a)])
                cur = pts[-1]
            elif name == "close":
                flush(closed=True)
        flush()
    return out


def table_parts(flags, r, c, n_rows, n_cols) -> list[str]:
    parts = ["wholeTbl"]
    header = 1 if flags["firstRow"] else 0
    footer = 1 if flags["lastRow"] else 0
    if flags["bandRow"] and header <= r < n_rows - footer:
        parts.append("band1H" if (r - header) % 2 == 0 else "band2H")
    if flags["bandCol"]:
        parts.append("band1V" if c % 2 == 0 else "band2V")
    if flags["firstCol"] and c == 0:
        parts.append("firstCol")
    if flags["lastCol"] and c == n_cols - 1:
        parts.append("lastCol")
    if flags["lastRow"] and r == n_rows - 1:
        parts.append("lastRow")
    if flags["firstRow"] and r == 0:
        parts.append("firstRow")
    return parts


def notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    try:
        frame = slide.notes_slide.notes_text_frame
        return frame.text.strip() if frame is not None else ""
    except Exception:
        return ""


# ---------------------------------------------------------------- animations and transitions

# PowerPoint's preset effects (presetClass, presetID) → the studio's (animate.js).
ENTRANCES = {1: "appear", 2: "flyIn", 3: "randomBars", 4: "shape", 5: "randomBars", 6: "shape", 8: "shape", 9: "fade", 10: "fade",
             12: "flyIn", 13: "shape", 14: "randomBars", 16: "split", 17: "stretch", 18: "wipe", 19: "swivel", 21: "wheel", 22: "wipe",
             23: "zoom", 26: "bounce", 30: "floatIn", 31: "growTurn", 35: "pinwheel", 37: "rise", 42: "floatIn", 45: "swivel", 47: "floatIn",
             49: "spinner", 53: "zoom", 55: "expand"}
EMPHASES = {1: "fillColor", 3: "fontColor", 6: "growShrink", 7: "lineColor", 8: "spin", 9: "transparency", 10: "boldFlash",
            14: "blink", 19: "desaturate", 21: "brighten", 26: "pulse", 27: "colorPulse", 32: "teeter", 34: "wave", 36: "pulse"}
FLY_DIRS = {1: "top", 2: "right", 3: "topRight", 4: "bottom", 6: "bottomRight", 8: "left", 9: "topLeft", 12: "bottomLeft"}
WIPE_DIRS = {1: "top", 2: "right", 4: "bottom", 8: "left"}


def timeline_of(slide_el, by_spid, deck: Deck) -> list[dict[str, Any]]:
    timing = slide_el.find(P + "timing")
    if timing is None:
        return []
    seq = None
    for node in timing.iter(P + "cTn"):
        if node.get("nodeType") == "mainSeq":
            seq = node
            break
    if seq is None:
        return []
    out = []
    for ctn in seq.iter(P + "cTn"):
        cls = ctn.get("presetClass")
        if not cls:
            continue
        target = next((t for t in ctn.iter(P + "spTgt")), None)
        spid = target.get("spid") if target is not None else None
        el = by_spid.get(spid or "")
        if not el:
            continue
        preset = int(ctn.get("presetID", 0))
        sub = int(ctn.get("presetSubtype", 0))
        start = {"clickEffect": "click", "withEffect": "with", "afterEffect": "after"}.get(ctn.get("nodeType"), "click")
        durs = [int(c.get("dur")) for c in ctn.iter(P + "cTn") if (c.get("dur") or "").isdigit()]
        delay_el = ctn.find(f"{P}stCondLst/{P}cond")
        delay = int(delay_el.get("delay")) if delay_el is not None and (delay_el.get("delay") or "").isdigit() else 0
        entry: dict[str, Any] = {"el": el, "start": start}
        if cls in ("entr", "exit"):
            fx = ENTRANCES.get(preset, "fade")
            entry.update({"cls": "in" if cls == "entr" else "out", "fx": fx})
            if fx == "flyIn" and FLY_DIRS.get(sub):
                entry["dir"] = FLY_DIRS[sub]
            if fx == "wipe" and WIPE_DIRS.get(sub):
                entry["dir"] = WIPE_DIRS[sub]
        elif cls == "emph":
            entry.update({"cls": "em", "fx": EMPHASES.get(preset, "pulse")})
        elif cls == "path":
            motion = next((m for m in ctn.iter(P + "animMotion")), None)
            pts = path_points(motion.get("path", "") if motion is not None else "", deck)
            if len(pts) < 2:
                continue
            entry.update({"cls": "path", "fx": "custom", "path": {"pts": pts}})
        else:
            continue
        if durs:
            entry["dur"] = max(1, max(durs))
        if delay:
            entry["delay"] = delay
        if target is not None and target.find(P + "txEl") is not None:
            entry["by"] = "para"
        out.append(entry)
    if len(out) > 5000:
        raise ValueError("アニメーションが5000件を超えています。")
    return out


def path_points(path: str, deck: Deck) -> list[list[float]]:
    """A motion path ("M 0 0 L 0.25 0.1 E", fractions of the slide) as pixel offsets."""
    nums = re.findall(r"-?\d*\.?\d+(?:[eE]-?\d+)?", path)
    pts = []
    for i in range(0, len(nums) - 1, 2):
        pts.append([round(float(nums[i]) * deck.cx * deck.k, 1), round(float(nums[i + 1]) * deck.cy * deck.k, 1)])
    if len(pts) > 2000:
        raise ValueError("アニメーションの軌跡が2000点を超えています。")
    return pts


TRANSITIONS = {"fade": "fade", "dissolve": "fade", "push": "push", "wipe": "wipe", "cover": "slide", "pull": "slide", "split": "curtain",
               "zoom": "zoom", "circle": "circle", "blinds": "blinds", "randomBar": "blinds", "checker": "blinds", "cut": "none",
               "wheel": "circle", "wedge": "circle", "diamond": "zoom", "plus": "zoom", "newsflash": "zoom", "strips": "wipe",
               "morph": "morph", "conveyor": "push", "pan": "push", "flip": "flip", "gallery": "flip", "vortex": "dive", "shred": "dive",
               "doors": "curtain", "window": "curtain", "reveal": "fade", "flash": "fade", "ripple": "circle", "honeycomb": "zoom",
               "glitter": "fade", "prism": "flip", "warp": "dive", "ferris": "flip", "flythrough": "dive", "switch": "flip", "curtains": "curtain"}


def transition_of(slide_el) -> dict[str, Any]:
    tr = slide_el.find(P + "transition")
    if tr is None:
        alt = slide_el.find(f"{{{NS['mc']}}}AlternateContent")
        if alt is not None:
            for branch in alt:
                tr = branch.find(P + "transition")
                if tr is not None:
                    break
    if tr is None:
        return {}
    out: dict[str, Any] = {}
    effect = next((c for c in tr if local(c) not in ("sndAc", "extLst")), None)
    kind = local(effect) if effect is not None else None
    if kind and TRANSITIONS.get(kind):
        out["transition"] = TRANSITIONS[kind]
        direction = transition_direction(out["transition"], effect)
        if direction:
            out["transitionDir"] = direction
    dur = tr.get(f"{{{NS['p14']}}}dur") or tr.get("dur")
    if dur and dur.isdigit():
        out["transitionDur"] = clamp(int(dur), 100, 10000)
    elif tr.get("spd"):
        out["transitionDur"] = {"slow": 1000, "med": 750, "fast": 500}.get(tr.get("spd"), 700)
    if tr.get("advTm", "").isdigit():
        out["advance"] = round(int(tr.get("advTm")) / 1000, 1)
    sound = transition_sound(tr)
    if sound:
        out["transitionSound"] = sound
    return out


# 効果のオプション: PowerPoint names the way a slide moves (l = to the left, so it comes from the right); the studio
# names where it comes from. Only the studio's other-than-default options are kept.
FROM_SIDE = {"l": "right", "r": "left", "u": "down", "d": "up"}
DEFAULT_DIR = {"slide": "right", "push": "down", "wipe": "right", "flip": "left", "blinds": "horizontal", "curtain": "vertical"}


def transition_direction(studio_kind: str, effect) -> str | None:
    """The studio's transitionDir for a PowerPoint transition element, or None for the default (or none)."""
    if studio_kind not in DEFAULT_DIR:
        return None
    if studio_kind in ("slide", "push", "wipe"):
        # Push and cover default to "l"; wipe to "l" as well (it sweeps in from the right).
        way = FROM_SIDE.get(effect.get("dir", "l"))
    elif studio_kind == "flip":
        way = "right" if effect.get("dir") == "r" else "left"
    elif studio_kind == "blinds":
        way = "vertical" if effect.get("dir") == "vert" else "horizontal"
    else:
        # Split: orient="vert" opens to the left and right (the studio's default curtain), "horz" up and down.
        way = "vertical" if effect.get("orient", "horz") == "vert" else "horizontal"
    return way if way and way != DEFAULT_DIR[studio_kind] else None


# PowerPoint's built-in transition sounds, matched to the studio's synthesised ones by their file name.
TRANSITION_SOUNDS = {
    "chime": "chime", "click": "click", "camera": "camera", "whoosh": "whoosh", "breeze": "whoosh", "wind": "whoosh",
    "arrow": "whoosh", "push": "whoosh", "suction": "whoosh", "drumroll": "drum", "hammer": "drum", "bomb": "drum",
    "explode": "drum", "applause": "applause", "coin": "coin", "cashreg": "coin", "laser": "bell", "voltage": "bell",
    "typewriter": "click",
}


def transition_sound(tr) -> str | None:
    """The slide's 画面切り替え sound: one of the studio's (by the built-in file's name), "stop", or none."""
    snd = tr.find(P + "sndAc")
    if snd is None:
        return None
    if snd.find(P + "endSnd") is not None:
        return "stop"
    st = snd.find(P + "stSnd")
    el = st.find(P + "snd") if st is not None else None
    if el is None:
        return None
    name = re.sub(r"\.[a-z0-9]+$", "", (el.get("name") or "").strip().lower())
    return TRANSITION_SOUNDS.get(name, "chime")


# ---------------------------------------------------------------- the whole deck

# ---------------------------------------------------------------- charts

C = "{%s}" % NS["c"]
MARKER_SYMBOLS = {"circle": "circle", "square": "square", "diamond": "diamond", "triangle": "triangle", "dash": "dash", "dot": "dot",
                  "x": "x", "plus": "plus", "star": "star", "none": "none", "auto": "circle", "picture": "square"}
LABEL_POSITIONS = {"outEnd", "inEnd", "ctr", "inBase", "t", "b", "l", "r", "bestFit"}
CHART_LABELS, CHART_SERIES = 500, 100
WEEKDAYS = "月火水木金土日"


def c_val(el, tag: str, default=None):
    node = el.find(C + tag) if el is not None else None
    return node.get("val", default) if node is not None else default


def c_bool(el, tag: str, default: bool = False) -> bool:
    """A chart on/off setting: <c:x/> alone means on."""
    node = el.find(C + tag) if el is not None else None
    if node is None:
        return default
    return node.get("val", "1") in ("1", "true")


def excel_date(serial: float, code: str) -> str:
    import datetime
    d = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(serial))

    def token(m):
        t = m.group(0)
        if t.startswith('"'):
            return t[1:-1]
        if t.startswith("\\"):
            return t[1:]
        if t.startswith("["):
            return ""
        low = t.lower()
        return {"yyyy": f"{d.year}", "yy": f"{d.year % 100:02d}", "mmmm": d.strftime("%B"), "mmm": d.strftime("%b"), "mm": f"{d.month:02d}",
                "m": f"{d.month}", "dddd": d.strftime("%A"), "ddd": d.strftime("%a"), "dd": f"{d.day:02d}", "d": f"{d.day}",
                "aaaa": WEEKDAYS[d.weekday()] + "曜日", "aaa": WEEKDAYS[d.weekday()]}.get(low, t)
    return re.sub(r'"[^"]*"|\\.|\[[^\]]*\]|yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|aaaa|aaa', token, code.split(";")[0], flags=re.I)


def number_text(value: float, code: str | None) -> str:
    """A category that is a number, as the chart shows it."""
    if code and re.search(r"[yd]|m{1,4}(?![^\[]*\])", re.sub(r'"[^"]*"|\[[^\]]*\]', "", code), re.I) and not re.search(r"[0#]", code):
        try:
            return excel_date(value, code)
        except (OverflowError, ValueError):
            pass
    if not code or code.lower() == "general":
        return str(int(value)) if float(value).is_integer() else f"{value:.10g}"
    section = code.split(";")[0]
    literal = re.findall(r'"([^"]*)"', section)
    bare = re.sub(r'"[^"]*"|\[[^\]]*\]', "", section)
    pct = "%" in bare
    m = re.search(r"[0#,]+(\.[0#]+)?", bare)
    dec = len(re.sub(r"[^0]", "", m.group(1))) if m and m.group(1) else 0
    v = Decimal(str(value * (100 if pct else 1))).quantize(Decimal(1).scaleb(-dec), rounding=ROUND_HALF_UP)  # Excel rounds halves up
    text = f"{v:,.{dec}f}" if m and "," in m.group(0) else f"{v:.{dec}f}"
    prefix = literal[0] if literal and section.find('"') < (section.find(m.group(0)) if m else 0) else ""
    suffix = "".join(literal[1:] if prefix else literal)
    return f"{prefix}{text}{'%' if pct else ''}{suffix}"


def cache_of(container):
    """A series' numbers or names from the chart's own cache: ([text or None], format code, numeric)."""
    if container is None:
        return [], None, False
    for path in ("numRef/numCache", "strRef/strCache", "numLit", "strLit", "multiLvlStrRef/multiLvlStrCache"):
        cache = container.find("/".join(C + t for t in path.split("/")))
        if cache is None:
            continue
        numeric = path.startswith("num")
        if local(cache) == "multiLvlStrCache":
            lvl = cache.find(C + "lvl")
            pts = lvl.findall(C + "pt") if lvl is not None else []
        else:
            pts = cache.findall(C + "pt")
        count = int(c_val(cache, "ptCount", "0") or 0)
        size = max([count] + [int(p.get("idx", 0)) + 1 for p in pts])
        out = [None] * min(size, 500)
        for pt in pts:
            i = int(pt.get("idx", 0))
            if i < len(out):
                out[i] = pt.findtext(C + "v")
        return out, cache.findtext(C + "formatCode"), numeric
    v = container.findtext(C + "v")
    return ([v] if v is not None else []), None, False


def _cx_levels(dim):
    """A chartEx dimension's levels, each the text at every point (the leaves first, then their parents)."""
    out = []
    for lvl in dim.findall(f"{{{NS['cx']}}}lvl"):
        pts = {int(pt.get("idx", 0)): (pt.text or "").strip() for pt in lvl.findall(f"{{{NS['cx']}}}pt")}
        count = max(int(lvl.get("ptCount") or 0), (max(pts) + 1) if pts else 0)
        out.append([pts.get(i, "") for i in range(min(count, 500))])
    return out


def _cx_num(text):
    try:
        v = float(str(text).replace(",", ""))
    except (TypeError, ValueError):
        return None
    return v if v == v and abs(v) < 1e12 else None


def read_chartex(root):
    """A chartEx part (cx:chartSpace) as the studio's chart: {type, labels, series, title?, opts?}; None for kinds the
    studio does not draw (a map, a pareto line…)."""
    q = lambda t: f"{{{NS['cx']}}}{t}"
    data = {}
    for d in root.iter(q("data")):
        cats, vals = [], []
        for dim in d:
            if local(dim) == "strDim" and dim.get("type") == "cat":
                cats = _cx_levels(dim)
            elif local(dim) == "numDim" and dim.get("type") in ("val", "size"):
                levels = _cx_levels(dim)
                vals = levels[0] if levels else []
        data[d.get("id")] = (cats, vals)
    series = [sr for sr in root.iter(q("series")) if not sr.get("hidden") == "1"]
    if not series:
        return None
    layout = series[0].get("layoutId", "")
    if any(sr.get("layoutId") == "paretoLine" for sr in series):
        series = [sr for sr in series if sr.get("layoutId") != "paretoLine"]

    def name_of(sr, i):
        v = sr.find(f"{q('tx')}/{q('txData')}/{q('v')}")
        return (v.text or "").strip() if v is not None and (v.text or "").strip() else f"系列{i + 1}"

    def data_of(sr):
        ref = sr.find(q("dataId"))
        return data.get(ref.get("val") if ref is not None else None, ([], []))

    out = None
    cats, vals = data_of(series[0])
    if layout in ("treemap", "sunburst"):
        leaves = cats[0] if cats else []
        parents = cats[1] if len(cats) > 1 else []
        labels, values, carry = [], [], ""
        for i, leaf in enumerate(leaves):
            parent = parents[i] if i < len(parents) and parents[i] else carry
            carry = parent
            labels.append(f"{parent}/{leaf}" if parent and leaf else (leaf or parent or f"項目{i + 1}"))
            values.append(_cx_num(vals[i]) if i < len(vals) else None)
        keep = [(l, v) for l, v in zip(labels, values) if v is not None and v > 0]
        if keep:
            out = {"type": layout, "labels": [l for l, _ in keep], "series": [{"name": name_of(series[0], 0), "values": [v for _, v in keep]}]}
    elif layout == "boxWhisker":
        groups = []
        for i, sr in enumerate(series):
            c, v = data_of(sr)
            nums = [_cx_num(x) for x in v]
            names = c[0] if c else []
            if len(series) == 1 and len({n for n in names if n}) > 1:
                for n in dict.fromkeys(names):
                    groups.append((n or f"系列{len(groups) + 1}", [x for x, m in zip(nums, names) if m == n and x is not None]))
            else:
                groups.append((name_of(sr, i), [x for x in nums if x is not None]))
        groups = [(n, v) for n, v in groups if v][:100]
        if groups:
            n = max(len(v) for _, v in groups)
            out = {"type": "boxplot", "labels": [str(i + 1) for i in range(n)], "series": [{"name": name, "values": v + [None] * (n - len(v))} for name, v in groups]}
    elif layout == "clusteredColumn":
        nums = [x for x in (_cx_num(v) for v in vals) if x is not None][:500]
        if nums:
            out = {"type": "histogram", "labels": [str(i + 1) for i in range(len(nums))], "series": [{"name": name_of(series[0], 0), "values": nums}]}
    elif layout in ("funnel", "waterfall"):
        labels = cats[0] if cats else [str(i + 1) for i in range(len(vals))]
        n = min(len(labels), len(vals))
        if n:
            out = {"type": layout, "labels": [labels[i] or f"項目{i + 1}" for i in range(n)], "series": [{"name": name_of(series[0], 0), "values": [_cx_num(vals[i]) or 0 for i in range(n)]}]}
            if layout == "waterfall":
                totals = [int(x.get("val")) for x in series[0].iter(q("idx")) if (x.get("val") or "").isdigit() and int(x.get("val")) < n]
                if totals:
                    out["opts"] = {"totals": sorted(set(totals))}
    if out:
        title = root.find(f"{q('chart')}/{q('title')}")
        if title is not None:
            text = " ".join(t.text or "" for t in title.iter() if local(t) in ("v", "t") and t.text).strip()
            if text:
                out["title"] = text[:80]
    return out


class ChartReader:
    """A chart part: its data (as the studio's chart) and its formatting (chart.style, drawn by objects.js officeChart)."""

    def __init__(self, sr: SlideReader, root):
        self.sr = sr
        self.deck = sr.deck
        self.root = root
        self.chart = root.find(C + "chart")
        tx1 = sr.theme.colors.get(sr.clr_map.get("tx1", "dk1"), "000000")
        try:
            rgb = tuple(int(tx1[i:i + 2], 16) for i in (0, 2, 4))
        except (TypeError, ValueError):
            rgb = (0, 0, 0)
        self.ink = "#" + "".join(f"{c:02x}" for c in _hsl_mod(rgb, 0.65, 0.35))

    # ---- formatting pieces

    def hex(self, container):
        c = self.sr.color(container)
        if not c:
            return None
        (r, g, b), alpha = c
        if alpha < 1:  # see-through colours over the white slide
            r, g, b = (round(v * alpha + 255 * (1 - alpha)) for v in (r, g, b))
        return f"#{r:02x}{g:02x}{b:02x}"

    def rpr_font(self, rpr) -> dict[str, Any]:
        out: dict[str, Any] = {}
        if rpr is None:
            return out
        if rpr.get("sz"):
            out["size"] = r2(self.deck.pt(int(rpr.get("sz")) / 100))
        if rpr.get("b") is not None:
            out["bold"] = rpr.get("b") in ("1", "true")
        if rpr.get("i") in ("1", "true"):
            out["italic"] = True
        fill = rpr.find(A + "solidFill")
        color = self.hex(fill) if fill is not None else None
        if color:
            out["color"] = color
        return out

    def font(self, el) -> dict[str, Any]:
        """The font a c:txPr (or a title's c:rich) gives: its first paragraph's defaults, then its first run."""
        if el is None:
            return {}
        txpr = el if local(el) in ("txPr", "rich") else el.find(C + "txPr")
        if txpr is None:
            return {}
        out = self.rpr_font(txpr.find(f".//{A}pPr/{A}defRPr"))
        out.update(self.rpr_font(txpr.find(f".//{A}r/{A}rPr")))
        return out

    def line(self, sppr, default=None):
        """'none', a colour, or the default when the chart does not say."""
        ln = sppr.find(A + "ln") if sppr is not None else None
        if ln is None:
            return default
        if ln.find(A + "noFill") is not None:
            return "none"
        fill = ln.find(A + "solidFill")
        return (self.hex(fill) if fill is not None else None) or default

    def fill(self, sppr):
        if sppr is None:
            return None
        if sppr.find(A + "noFill") is not None:
            return "#ffffff"
        f = self.sr.fill_of(sppr)
        if isinstance(f, tuple) and f[0] == "solid":
            (r, g, b), alpha = f[1]
            if alpha < 1:
                r, g, b = (round(v * alpha + 255 * (1 - alpha)) for v in (r, g, b))
            return f"#{r:02x}{g:02x}{b:02x}"
        return None

    def accent(self, i: int) -> str:
        name = f"accent{i % 6 + 1}"
        return "#" + self.sr.theme.colors.get(name, "4472C4").lower()

    def label_flags(self, el, src_fmt, base=None) -> dict[str, Any]:
        out: dict[str, Any] = {}
        for tag, key in (("showVal", "val"), ("showPercent", "pct"), ("showCatName", "cat"), ("showSerName", "ser")):
            node = el.find(C + tag)
            on = (node.get("val", "1") in ("1", "true")) if node is not None else bool(base and base.get(key))
            if on:
                out[key] = True
        num_fmt = el.find(C + "numFmt")
        if num_fmt is not None and num_fmt.get("sourceLinked") != "1" and num_fmt.get("formatCode"):
            out["format"] = num_fmt.get("formatCode")
        elif base and base.get("format"):
            out["format"] = base["format"]
        elif src_fmt and src_fmt.lower() != "general":
            out["format"] = src_fmt
        pos = c_val(el, "dLblPos")
        if pos in LABEL_POSITIONS:
            out["pos"] = pos
        elif base and base.get("pos"):
            out["pos"] = base["pos"]
        f = self.font(el)
        if f:
            out["font"] = f
        return out

    def data_labels(self, ser, group, src_fmt, name, labels, values):
        dlbls = ser.find(C + "dLbls")
        if dlbls is None:
            dlbls = group.find(C + "dLbls")
        if dlbls is None or c_bool(dlbls, "delete"):
            return None, []
        base = self.label_flags(dlbls, src_fmt)
        shown = any(base.get(k) for k in ("val", "pct", "cat", "ser"))
        points = []
        for d in dlbls.findall(C + "dLbl"):
            i = int(c_val(d, "idx", "0") or 0)
            if i >= CHART_LABELS:
                continue
            if c_bool(d, "delete"):
                points.append({"i": i, "show": False})
                continue
            pl = self.label_flags(d, src_fmt, base)
            rich = d.find(f"{C}tx/{C}rich")
            if rich is not None:
                runs = []
                for para_i, para in enumerate(rich.findall(A + "p")):
                    if para_i:
                        runs.append({"t": "\n"})
                    for r in para:
                        kind = local(r)
                        if kind == "r":
                            text = r.findtext(A + "t") or ""
                        elif kind == "fld":
                            ftype = r.get("type", "")
                            text = (name if ftype == "SERIESNAME" else labels[i] if ftype == "CATEGORYNAME" and i < len(labels)
                                    else number_text(values[i], pl.get("format")) if ftype == "VALUE" and i < len(values) else r.findtext(A + "t") or "")
                        else:
                            continue
                        if text:
                            f = self.rpr_font(r.find(A + "rPr"))
                            runs.append({"t": text, **({"font": f} if f else {})})
                if runs:
                    pl["runs"] = [r for r in runs if r["t"] != "\n"] if len(rich.findall(A + "p")) == 1 else runs
            if pl.get("runs") or any(pl.get(k) for k in ("val", "pct", "cat", "ser")):
                points.append({"i": i, **pl})
            else:
                points.append({"i": i, "show": False})
        return (base if shown else None), points

    def axis(self, ax, src_fmt, *, is_val: bool) -> dict[str, Any]:
        if ax is None:
            return {"hide": True, "line": "none"}
        out: dict[str, Any] = {}
        deleted = c_bool(ax, "delete")
        if deleted or c_val(ax, "tickLblPos") == "none":
            out["hide"] = True
        if c_val(ax.find(C + "scaling"), "orientation") == "maxMin":
            out["reverse"] = True
        if is_val:
            for tag in ("min", "max"):
                v = c_val(ax.find(C + "scaling"), tag)
                if v is not None:
                    out[tag] = float(v)
            if c_val(ax, "majorUnit"):
                out["step"] = float(c_val(ax, "majorUnit"))
            num_fmt = ax.find(C + "numFmt")
            if num_fmt is not None and num_fmt.get("sourceLinked") != "1" and num_fmt.get("formatCode", "General").lower() != "general":
                out["format"] = num_fmt.get("formatCode")
            elif src_fmt and src_fmt.lower() != "general":
                out["format"] = src_fmt
        out["line"] = "none" if deleted else self.line(ax.find(C + "spPr"), "#868686")
        grid = ax.find(C + "majorGridlines")
        if grid is not None:
            out["grid"] = self.line(grid.find(C + "spPr"), "#d9d9d9")
        f = self.font(ax)
        if f:
            out["font"] = f
        return out

    # ---- the chart

    def read(self) -> dict[str, Any] | None:
        chart = self.chart
        plot = chart.find(C + "plotArea") if chart is not None else None
        if plot is None:
            return None
        groups = [g for g in plot if local(g).endswith("Chart")]
        if not groups:
            return None
        if sum(len(g.findall(C + "ser")) for g in groups) > CHART_SERIES:
            raise ValueError(f"スライド{self.sr.number}のグラフが100系列を超えています。")
        if any(int(node.get("val", 0)) > CHART_LABELS for node in plot.iter(C + "ptCount")):
            raise ValueError(f"スライド{self.sr.number}のグラフが500項目を超えています。")
        base_font = {"size": r2(self.deck.pt(18)), "color": self.ink}
        base_font.update(self.font(self.root.find(C + "txPr")))
        base_font.update({k: v for k, v in self.font(chart.find(C + "txPr")).items()})
        series: list[dict[str, Any]] = []
        styles: list[dict[str, Any]] = []
        labels: list[str] = []
        kinds: list[str] = []
        st: dict[str, Any] = {"font": base_font}
        first_bar = None
        pie = None
        for g in groups:
            gname = local(g)
            if gname in ("barChart", "bar3DChart"):
                kind = "bar"
                if first_bar is None:
                    first_bar = g
            elif gname in ("pieChart", "pie3DChart", "ofPieChart", "doughnutChart"):
                kind = "pie"
                pie = g
            elif gname in ("areaChart", "area3DChart"):
                kind = "area"
            else:
                kind = "line"
            vary = c_bool(g, "varyColors", default=kind == "pie")
            group_markers = c_bool(g, "marker", default=True)
            sers = sorted(g.findall(C + "ser"), key=lambda e: int(c_val(e, "order", "0") or 0))
            for ser in sers:
                if len(series) >= CHART_SERIES:
                    break
                name_vals, _, _ = cache_of(ser.find(C + "tx"))
                name = (name_vals[0] if name_vals and name_vals[0] else f"系列{len(series) + 1}")[:500]
                cat_el = ser.find(C + "cat")
                if cat_el is None:
                    cat_el = ser.find(C + "xVal")
                val_el = ser.find(C + "val")
                if val_el is None:
                    val_el = ser.find(C + "yVal")
                cats, cat_fmt, cat_num = cache_of(cat_el)
                vals, val_fmt, _ = cache_of(val_el)
                if not labels and cats:
                    labels = [(number_text(float(c), cat_fmt) if cat_num and c not in (None, "") and re.match(r"^-?[\d.]+(e-?\d+)?$", c, re.I) else str(c or ""))[:500] for c in cats][:CHART_LABELS]
                values = []
                for v in vals[:CHART_LABELS]:
                    try:
                        values.append(float(v))
                    except (TypeError, ValueError):
                        values.append(0.0)
                sppr = ser.find(C + "spPr")
                index = int(c_val(ser, "idx", str(len(series))) or 0)
                one: dict[str, Any] = {"kind": kind if kind != "pie" else "bar"}
                if kind == "line":
                    color = self.line(sppr) if sppr is not None else None
                    one["color"] = color if color and color != "none" else self.accent(index)
                    ln = sppr.find(A + "ln") if sppr is not None else None
                    if ln is not None and ln.find(A + "noFill") is not None:
                        one["width"] = 0
                    else:
                        one["width"] = r2(self.deck.px(int(ln.get("w")) if ln is not None and ln.get("w") else 28575))
                    dash = ln.find(A + "prstDash") if ln is not None else None
                    if dash is not None and DASHES.get(dash.get("val")):
                        one["dash"] = DASHES[dash.get("val")]
                    if c_bool(ser, "smooth"):
                        one["smooth"] = True
                    mk = ser.find(C + "marker")
                    symbol = MARKER_SYMBOLS.get(c_val(mk, "symbol", "auto") if mk is not None else ("auto" if group_markers else "none"), "circle")
                    marker: dict[str, Any] = {"s": symbol}
                    if mk is not None and c_val(mk, "size"):
                        marker["z"] = r2(self.deck.pt(float(c_val(mk, "size"))))
                    if mk is not None:
                        mcolor = self.fill(mk.find(C + "spPr"))
                        if mcolor:
                            marker["color"] = mcolor
                    at = []
                    for dpt in ser.findall(C + "dPt"):
                        dmk = dpt.find(C + "marker")
                        if dmk is not None and c_val(dmk, "symbol", "auto") != "none":
                            at.append(int(c_val(dpt, "idx", "0") or 0))
                            if c_val(dmk, "size") and "z" not in marker:
                                marker["z"] = r2(self.deck.pt(float(c_val(dmk, "size"))))
                            dcolor = self.fill(dmk.find(C + "spPr"))
                            if dcolor and "color" not in marker:
                                marker["color"] = dcolor
                    if at and symbol == "none":
                        marker["at"] = at
                    one["marker"] = marker
                else:
                    color = self.fill(sppr)
                    one["color"] = color or self.accent(index)
                    points = []
                    if vary and (kind == "pie" or len(sers) == 1):
                        points = [{"i": i, "color": self.accent(i)} for i in range(min(len(values), CHART_LABELS))]
                    for dpt in ser.findall(C + "dPt"):
                        i = int(c_val(dpt, "idx", "0") or 0)
                        pc = self.fill(dpt.find(C + "spPr"))
                        if pc and i < CHART_LABELS:
                            points = [p for p in points if p["i"] != i] + [{"i": i, "color": pc}]
                    if points:
                        one["points"] = sorted(points, key=lambda p: p["i"])
                base, point_labels = self.data_labels(ser, g, val_fmt, name, labels, values)
                if base:
                    one["label"] = base
                if point_labels:
                    one["pointLabels"] = point_labels
                series.append({"name": name, "values": values})
                styles.append(one)
                kinds.append(kind)
            if pie is not None:
                break
        if not series:
            return None
        n = max(len(labels), max(len(s["values"]) for s in series))
        if not labels:
            labels = [str(i + 1) for i in range(n)]
        labels = labels[:CHART_LABELS]
        for s_ in series:
            s_["values"] = (s_["values"] + [0.0] * len(labels))[:len(labels)]
        # The studio's kind of chart (its editor offers these), and how it is drawn
        if pie is not None:
            ctype = "donut"
            st["hole"] = float(c_val(pie, "holeSize", "10") or 10) if local(pie) == "doughnutChart" else 0
            if c_val(pie, "firstSliceAng"):
                st["angle"] = float(c_val(pie, "firstSliceAng"))
            series, styles = series[:1], styles[:1]
        elif first_bar is not None:
            grouping = c_val(first_bar, "grouping", "clustered")
            bars = kinds.count("bar")
            if len(set(kinds)) > 1:
                ctype = "combo"
            elif grouping == "percentStacked":
                ctype = "100-stacked-bar"
            elif grouping == "stacked":
                ctype = "stacked-bar"
            else:
                ctype = "clustered-bar" if bars > 1 else "bar"
            if grouping in ("stacked", "percentStacked"):
                st["stack"] = "percent" if grouping == "percentStacked" else "stacked"
            st["dir"] = "bar" if c_val(first_bar, "barDir", "col") == "bar" else "col"
            st["gap"] = float(c_val(first_bar, "gapWidth", "150") or 150)
            if c_val(first_bar, "overlap"):
                st["overlap"] = float(c_val(first_bar, "overlap"))
            if ctype == "combo" and kinds[0] != "bar":  # the editor's combo is bars first, then a line
                order = sorted(range(len(series)), key=lambda i: kinds[i] != "bar")
                series, styles = [series[i] for i in order], [styles[i] for i in order]
        else:
            ctype = "line" if len(series) == 1 else "multi-line"
        st["type"] = ctype
        st["series"] = styles
        # Axes
        cat_ax = plot.find(C + "catAx")
        if cat_ax is None:
            cat_ax = plot.find(C + "dateAx")
        val_axes = plot.findall(C + "valAx")
        val_ax = val_axes[0] if val_axes else None
        if cat_ax is None and len(val_axes) > 1:  # scatter: x is a value axis too
            cat_ax, val_ax = val_axes[0], val_axes[1]
        if pie is None:
            src_fmt = None
            for g in groups:
                ser = g.find(C + "ser")
                if ser is not None:
                    src_fmt = cache_of(ser.find(C + "val"))[1]
                    break
            st["cat"] = self.axis(cat_ax, None, is_val=False)
            st["val"] = self.axis(val_ax, src_fmt, is_val=True)
            if val_ax is not None and c_val(val_ax, "crossBetween") == "midCat":
                st["cat"]["edge"] = True
        # Legend, title and the plot area's place
        legend = chart.find(C + "legend")
        if legend is not None:
            pos = c_val(legend, "legendPos", "r")
            st["legend"] = {"pos": pos if pos in ("b", "t", "r", "l", "tr") else "r"}
            f = self.font(legend)
            if f:
                st["legend"]["font"] = f
        title = ""
        title_el = chart.find(C + "title")
        if title_el is not None and not c_bool(chart, "autoTitleDeleted"):
            rich = title_el.find(f"{C}tx/{C}rich")
            if rich is not None:
                title = "\n".join("".join(t.text or "" for t in p.iter(A + "t")) for p in rich.findall(A + "p")).strip()
                f = self.font(rich)
                if f:
                    st["title"] = f
            elif len(series) == 1:
                title = series[0]["name"]
            if not st.get("title"):
                f = self.font(title_el)
                if f:
                    st["title"] = f
        manual = plot.find(f"{C}layout/{C}manualLayout")
        if manual is not None and c_val(manual, "xMode", "factor") == "edge" and c_val(manual, "yMode", "factor") == "edge":
            try:
                x, y = float(c_val(manual, "x", "0")), float(c_val(manual, "y", "0"))
                w, h = float(c_val(manual, "w", "1")), float(c_val(manual, "h", "1"))
                if c_val(manual, "wMode") == "edge":
                    w -= x
                if c_val(manual, "hMode") == "edge":
                    h -= y
                st["plot"] = {"x": r2(x), "y": r2(y), "w": r2(max(0.05, w)), "h": r2(max(0.05, h)), **({"outer": True} if c_val(manual, "layoutTarget", "outer") != "inner" else {})}
            except (TypeError, ValueError):
                pass
        out: dict[str, Any] = {"type": ctype, "labels": labels, "series": series, "style": st}
        if title:
            out["title"] = title[:80]
        return out


def read_pptx_exact(data: bytes) -> dict[str, Any]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    if len(prs.slides) > 500:
        raise ValueError("スライドが500枚を超えています。")
    deck = Deck(prs, {})
    # Each slide that comes over keeps an id, so links between slides still go where they went.
    shown = list(prs.slides)
    deck.sids = {str(s.part.partname): f"p{i + 1}" for i, s in enumerate(shown)}
    slides = []
    number = 0
    for index, slide in enumerate(prs.slides):
        number += 1
        reader = SlideReader(deck, slide, len(slides), number)
        imported = reader.read()
        if slide._element.get("show") in ("0", "false"):
            imported["hidden"] = True
            deck.stats["hidden"] += 1
        slides.append(imported)
    deck.stats["slides"] = len(slides)
    title = (prs.core_properties.title or "").strip() or (slides[0]["title"] if slides else "") or "取り込んだ資料"
    return {"deckTitle": title[:100], "slideData": slides, "stats": deck.stats, "fidelity": "exact", "sej": deck.sej,
            "size": {"w": deck.cx, "h": deck.cy}}
