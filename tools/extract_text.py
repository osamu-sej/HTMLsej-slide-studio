#!/usr/bin/env python3
"""Extract plain text from PPTX, DOCX or PDF bytes on stdin for use as source material."""

from __future__ import annotations

import io
import json
import re
import sys
import zipfile

from lxml import etree

MAX_CHARS = 60_000


def from_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for index, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if lines or notes:
            parts.append(f"[スライド{index}]\n" + "\n".join(lines) + (f"\n（ノート）{notes}" if notes else ""))
    return "\n\n".join(parts)


def from_docx(data: bytes) -> str:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    lines = []
    body = root.find("w:body", ns)
    for block in body if body is not None else []:
        tag = etree.QName(block).localname
        if tag == "p":
            text = "".join(block.xpath(".//w:t/text()", namespaces=ns)).strip()
            if text:
                lines.append(text)
        elif tag == "tbl":
            for row in block.xpath(".//w:tr", namespaces=ns):
                cells = ["".join(cell.xpath(".//w:t/text()", namespaces=ns)).strip() for cell in row.xpath("./w:tc", namespaces=ns)]
                if any(cells):
                    lines.append(" | ".join(cells))
    return "\n".join(lines)


def from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError("PDFの読み込みに必要なライブラリがありません") from error
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[ページ{index}]\n{text}")
    return "\n\n".join(pages)


def main() -> int:
    name = (sys.argv[1] if len(sys.argv) > 1 else "").lower()
    data = sys.stdin.buffer.read()
    if name.endswith(".pptx"):
        text = from_pptx(data)
    elif name.endswith(".docx"):
        text = from_docx(data)
    elif name.endswith(".pdf"):
        text = from_pdf(data)
    else:
        raise ValueError("対応していないファイル形式です（pptx / docx / pdf）")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("文章を取り出せませんでした（画像だけのファイルの可能性があります）")
    print(json.dumps({"text": text[:MAX_CHARS], "truncated": len(text) > MAX_CHARS}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
